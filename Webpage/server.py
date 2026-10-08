# server.py
import os
import sys
import json
import re
import subprocess
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from flask import Flask, request, render_template_string, jsonify, abort, redirect, url_for, send_file, Response
import pandas as pd
import numpy as np
from sqlalchemy import text
from openpyxl import load_workbook

from ui import (
    ERR_TPL,
    INDEX_TPL,
    SUBPAGE_TPL,
    ITEM_TPL,
    ITEM_INFO_TPL,
    QUOTE_TPL,
    PERIPHERAL_STATUS_TPL,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
MRP_MODULE_DIR = REPO_ROOT / "MRP_System 3.0"
if str(MRP_MODULE_DIR) not in sys.path:
    sys.path.insert(0, str(MRP_MODULE_DIR))

from mrp_system.normalize.mrp_normalize import normalize_item
from mrp_system.ledger.atp import build_atp_view, earliest_atp_strict, earliest_atp_for_items_strict
from mrp_system.runtime.db_config import get_engine, DATABASE_DSN
from mrp_system.quotation_cards import load_item_cards
from mrp_system.purchase_order_lookup import purchase_orders
from mrp_system.runtime.constants import PLACEHOLDER_DATE
from mrp_system.runtime.paths import PERIPHERAL_STATUS_FILE

app = Flask(__name__)
app.register_blueprint(purchase_orders)

# =========================
# DB ENGINE
# =========================
engine = get_engine()

# =========================
# Data cache
# =========================
SO_INV: pd.DataFrame | None = None
INVENTORY_STATUS: pd.DataFrame | None = None
SAP: pd.DataFrame | None = None
OPEN_PO: pd.DataFrame | None = None
LEDGER: pd.DataFrame | None = None
ITEM_ATP: pd.DataFrame | None = None
RECEIVING_LOG: pd.DataFrame | None = None
ITEM_INFO: pd.DataFrame | None = None
_LAST_LOAD_ERR: str | None = None
_LAST_LOADED_AT: datetime | None = None
ITEM_SUGGEST_CACHE: list[str] = []
GLOBAL_SEARCH_INDEX: list[dict[str, str]] = []
ITEM_INFO_SUGGEST_CACHE: list[str] = []
ITEM_INFO_PHOTO_BY_SUGGESTION: dict[str, bool] = {}
QUOTE_ITEM_SUGGEST_ROWS: list[dict[str, object]] = []
SO_LOOKUP_BASE: pd.DataFrame | None = None
WAITING_ITEMS_BY_QB: dict[str, str] = {}
LEDGER_ITEM_INDEX: dict[str, pd.DataFrame] = {}
PDF_DB_SEARCH_CACHE: dict[tuple[str, int], list[dict]] = {}
INDEX_VIEW_CACHE: dict[tuple[str, str], dict] = {}
QUOTATION_VIEW_CACHE: dict[tuple[str, int], dict] = {}
PERIPHERAL_STATUS_CACHE: dict | None = None
PERIPHERAL_STATUS_CACHE_KEY: tuple[str, int, int] | None = None
READY_ASSIGN_CACHE: list[dict] | None = None
RECENT_HOME_SEARCHES: list[dict[str, str]] = []

# =========================
# PDF settings/cache
# =========================
# Configure a root folder that contains PDF files named by order id (e.g. SO-12345.pdf)
PDF_FOLDER = os.getenv("PDF_FOLDER", "")
PDF_VIEW_BASE_URL = os.getenv("PDF_VIEW_BASE_URL", "http://192.168.60.215:5001/view_file")
# Map of order_id (stem of filename) -> {file_name, file_path}
PDF_MAP: dict[str, dict[str, str]] = {}

TABLE_HEADER_LABELS = {
    "Item": "Item",
    "Qty(-)": "Qty (-)",
    "Available": "Available",
    "Available + Pre-installed PO": "Avail + Pre-PO",
    "On Hand": "On Hand",
    "On Sales Order": "On SO",
    "On PO": "On PO",
    "Assigned Q'ty": "Assigned Qty",
    "On Hand - WIP": "On Hand - WIP",
    "Available + On PO": "Avail + On PO",
    "Sales/Week": "Sales / Week",
    "Recommended Restock Qty": "Restock Qty",
    "Component_Status": "Status",
    "Ship Date": "Ship Date",
}

# -------- helpers --------
def _safe_date_col(df: pd.DataFrame, col: str):
    if col in df.columns:
        df[col] = pd.to_datetime(df[col], errors="coerce")

def _to_date_str(s: pd.Series, fmt="%Y-%m-%d") -> pd.Series:
    s = pd.to_datetime(s, errors="coerce")
    return s.apply(lambda x: x.strftime(fmt) if pd.notnull(x) else "")


def _format_num(value, digits: int = 1) -> str:
    try:
        num = float(value)
    except Exception:
        return ""
    if not np.isfinite(num):
        return ""
    if abs(num - round(num)) < 0.000001:
        return str(int(round(num)))
    return f"{num:.{digits}f}".rstrip("0").rstrip(".")


def _normalize_so_key(v: str) -> str:
    """Normalize SO/QB key for tolerant matching (SO-20260328 == so20260328)."""
    s = str(v or "").upper().strip()
    return re.sub(r"[^A-Z0-9]", "", s)


def _read_table(schema: str, table: str) -> pd.DataFrame:
    sql = f'SELECT * FROM "{schema}"."{table}"'
    return pd.read_sql_query(text(sql), con=engine)


def _pdf_view_url_for_path(path: str | None) -> str | None:
    raw = str(path or "").strip()
    if not raw or not PDF_VIEW_BASE_URL:
        return None
    return f"{PDF_VIEW_BASE_URL.rstrip('/')}/{quote(raw, safe='')}"


def _load_item_info(force: bool = False) -> pd.DataFrame:
    global ITEM_INFO, ITEM_INFO_SUGGEST_CACHE, ITEM_INFO_PHOTO_BY_SUGGESTION
    if force or ITEM_INFO is None:
        item_info = _read_table("public", "Item Info")
        ITEM_INFO = item_info

        suggest_items: list[str] = []
        photo_by_suggestion: dict[str, bool] = {}
        for _, row in item_info.iterrows():
            has_photo = bool(str(row.get("Photo File", "") or "").strip())
            for col in ("Name", "Part Name"):
                if col not in item_info.columns:
                    continue
                value = str(row.get(col, "") or "").strip()
                if not value:
                    continue
                suggest_items.append(value)
                photo_by_suggestion[value] = photo_by_suggestion.get(value, False) or has_photo
        ITEM_INFO_SUGGEST_CACHE = sorted(set(suggest_items))
        ITEM_INFO_PHOTO_BY_SUGGESTION = photo_by_suggestion
    return ITEM_INFO


def _resolve_item_photo_path(photo_file: str) -> Path | None:
    raw = str(photo_file or "").strip()
    if not raw:
        return None

    allowed_root = (Path.home() / "OneDrive - neousys-tech" / "Share NTA Warehouse" / "Product List").resolve()
    raw_path = Path(raw)
    candidates: list[Path] = []

    if raw_path.is_absolute():
        candidates.append(raw_path)
    else:
        candidates.append(Path.home() / raw_path)
        candidates.append(allowed_root / raw_path.name)

    expanded: list[Path] = []
    known_photo_suffixes = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf"}
    for candidate in candidates:
        expanded.append(candidate)
        if candidate.suffix.lower() not in known_photo_suffixes:
            for suffix in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf"):
                expanded.append(candidate.with_suffix(suffix))
                expanded.append(Path(str(candidate) + suffix))

    for candidate in expanded:
        try:
            resolved = candidate.resolve()
            resolved.relative_to(allowed_root)
        except Exception:
            continue
        if resolved.is_file():
            return resolved
    return None


def _hydrate_onedrive_file(path: Path) -> None:
    try:
        with path.open("rb") as handle:
            handle.read(1)
        return
    except OSError:
        if os.name != "nt":
            raise

    subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-Command",
            "& { param($p) Get-Content -LiteralPath $p -Encoding Byte -TotalCount 1 | Out-Null }",
            str(path),
        ],
        check=True,
        timeout=60,
    )


# ---------- PDF DB helpers (no Flask-SQLAlchemy) ----------
def _pdf_db_search_by_filename(search_query: str, limit: int = 10) -> list[dict]:
    """Search pdf_file_log by file_name ILIKE %search_query% using SQLAlchemy Core.
    Returns list of dict rows; empty if table missing or error.
    """
    if not search_query:
        return []
    key = (str(search_query).strip().lower(), int(limit))
    if key in PDF_DB_SEARCH_CACHE:
        return PDF_DB_SEARCH_CACHE[key]
    try:
        sql = text(
            'SELECT id, order_id, file_name, file_path, extracted_data '
            'FROM "public"."pdf_file_log" WHERE file_name ILIKE :q '
            'ORDER BY id DESC LIMIT :lim'
        )
        with engine.connect() as conn:
            res = conn.execute(sql, {"q": f"%{search_query}%", "lim": limit})
            out = [dict(row) for row in res.mappings().all()]
            if len(PDF_DB_SEARCH_CACHE) >= 512:
                PDF_DB_SEARCH_CACHE.clear()
            PDF_DB_SEARCH_CACHE[key] = out
            return out
    except Exception:
        # Table might not exist, or permissions issues; return empty silently
        return []

def _pdf_db_get_by_id(pdf_id: int) -> dict | None:
    try:
        sql = text(
            'SELECT id, order_id, file_name, file_path, extracted_data '
            'FROM "public"."pdf_file_log" WHERE id = :id'
        )
        with engine.connect() as conn:
            res = conn.execute(sql, {"id": pdf_id})
            row = res.mappings().first()
            return dict(row) if row else None
    except Exception:
        return None


def _validate_paths(paths: list[str]):
    for p in paths:
        if not os.path.exists(p):
            # Using print to avoid coupling to any logger; environment prints to console
            print(f"[pdf] Path does not exist: {p}")
        else:
            print(f"[pdf] Valid path: {p}")

def _scan_pdf_folder(folder_path: str) -> dict[str, dict[str, str]]:
    data: dict[str, dict[str, str]] = {}
    if not folder_path:
        return data
    if not os.path.isdir(folder_path):
        print(f"[pdf] PDF_FOLDER is not a directory: {folder_path}")
        return data
    for root, _dirs, files in os.walk(folder_path):
        for name in files:
            if name.lower().endswith(".pdf"):
                path = os.path.join(root, name)
                order_id = os.path.splitext(name)[0]
                data[order_id.upper()] = {"file_name": name, "file_path": path}
    print(f"[pdf] scanned {len(data)} PDF(s) under {folder_path}")
    return data

def _load_pdf_map(force: bool = False):
    global PDF_MAP
    if PDF_MAP and not force:
        return
    if PDF_FOLDER:
        _validate_paths([PDF_FOLDER])
        PDF_MAP = _scan_pdf_folder(PDF_FOLDER)
    else:
        PDF_MAP = {}

def _build_runtime_indexes(
    so_src: pd.DataFrame, ledger_src: pd.DataFrame
) -> tuple[pd.DataFrame, dict[str, str], dict[str, pd.DataFrame]]:
    """Build read-optimized indexes used by index() and quotation_lookup()."""
    so = so_src.copy()
    if "QB Num" not in so.columns:
        so["QB Num"] = ""
    if "Name" not in so.columns:
        so["Name"] = ""
    so["__qb_upper"] = so["QB Num"].astype(str).str.upper().str.strip()
    so["__name_text"] = so["Name"].astype(str).str.strip()

    waiting_map: dict[str, str] = {}
    if {"QB Num", "Item", "Component_Status"}.issubset(so.columns):
        so_wait = so.loc[so["Component_Status"].isin(["Waiting", "Shortage"])].copy()
        so_wait["QB Num"] = so_wait["QB Num"].astype(str).str.strip()
        so_wait["Item"] = so_wait["Item"].astype(str).str.strip()
        if not so_wait.empty:
            waiting_map = (
                so_wait.groupby("QB Num")["Item"]
                .apply(lambda s: "{%s}" % ", ".join(sorted(set(i for i in s if i))))
                .to_dict()
            )

    ledger_index: dict[str, pd.DataFrame] = {}
    if ledger_src is not None and not ledger_src.empty:
        item_col = "Item" if "Item" in ledger_src.columns else ("Item_raw" if "Item_raw" in ledger_src.columns else None)
        if item_col is not None:
            work = ledger_src.copy()
            work[item_col] = work[item_col].astype(str).str.strip()
            work["__item_lookup__"] = work[item_col].str.upper()
            for item_key, grp in work.groupby("__item_lookup__", sort=False):
                ledger_index[str(item_key)] = grp.drop(columns=["__item_lookup__"]).copy()

    return so, waiting_map, ledger_index


def _build_quote_item_summaries(
    inventory_src: pd.DataFrame,
    ledger_src: pd.DataFrame,
) -> list[dict[str, object]]:
    if inventory_src is None or inventory_src.empty or "Part_Number" not in inventory_src.columns:
        return []

    inv = inventory_src.copy()
    inv["Part_Number"] = inv["Part_Number"].astype(str).str.strip()
    inv = inv.loc[inv["Part_Number"].ne("")].copy()
    if "On Hand" not in inv.columns:
        inv["On Hand"] = pd.NA
    inv["On Hand"] = pd.to_numeric(inv["On Hand"], errors="coerce")
    if "Available" not in inv.columns:
        inv["Available"] = 0.0
    if "Max" not in inv.columns:
        inv["Max"] = pd.NA
    inv["Available"] = pd.to_numeric(inv["Available"], errors="coerce")
    inv["Max"] = pd.to_numeric(inv["Max"], errors="coerce")
    inv_summary = (
        inv.groupby("Part_Number", as_index=False)[["On Hand", "Available", "Max"]]
        .first()
        .rename(columns={"Part_Number": "item", "On Hand": "on_hand", "Available": "available", "Max": "max_flag"})
    )

    records = inv_summary.sort_values("item", kind="mergesort").to_dict(orient="records")
    cleaned: list[dict[str, object]] = []
    for rec in records:
        out: dict[str, object] = {}
        for key, val in rec.items():
            out[key] = None if pd.isna(val) else val
        max_flag = out.get("max_flag")
        if max_flag == 0:
            out["highlight"] = "red"
        elif max_flag == 99:
            out["highlight"] = "green"
        else:
            out["highlight"] = ""
        cleaned.append(out)
    return cleaned

def _build_global_search_index(so: pd.DataFrame, inventory: pd.DataFrame) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    seen: set[tuple[str, str]] = set()

    customer_by_so: dict[str, str] = {}
    if so is not None and not so.empty and {"QB Num", "Name"}.issubset(so.columns):
        for qb_num, group in so.groupby("QB Num", dropna=True, sort=False):
            names = group["Name"].dropna().astype(str).str.strip()
            names = names.loc[names.ne("")]
            if not names.empty:
                customer_by_so[str(qb_num).strip().casefold()] = names.iloc[0]

    ship_date_by_so: dict[str, str] = {}
    if so is not None and not so.empty and {"QB Num", "Ship Date"}.issubset(so.columns):
        for qb_num, group in so.groupby("QB Num", dropna=True, sort=False):
            ship_dates = pd.to_datetime(group["Ship Date"], errors="coerce").dropna()
            if not ship_dates.empty:
                ship_date_by_so[str(qb_num).strip().casefold()] = ship_dates.min().strftime("%Y-%m-%d")

    on_hand_by_item: dict[str, str] = {}
    if inventory is not None and not inventory.empty and {"Part_Number", "On Hand"}.issubset(inventory.columns):
        for item, group in inventory.groupby("Part_Number", dropna=True):
            item_key = str(item or "").strip().casefold()
            if item_key:
                on_hand_by_item[item_key] = _format_num(_aggregate_metric(group["On Hand"]))

    def add(
        kind: str,
        label: object,
        href: str,
        *,
        on_hand: str = "",
        customer: str = "",
        ship_date: str = "",
    ) -> None:
        value = str(label or "").strip()
        if not value:
            return
        key = (kind, value.upper())
        if key in seen:
            return
        seen.add(key)
        entry: dict[str, object] = {"type": kind, "label": value, "href": href, "search": value.lower()}
        if kind == "Item":
            entry["on_hand"] = on_hand
        if kind == "SO":
            entry["customer"] = customer
            entry["ship_date"] = ship_date
        entries.append(entry)

    if so is not None and not so.empty:
        if "QB Num" in so.columns:
            for value in so["QB Num"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist():
                add(
                    "SO",
                    value,
                    "/?so=" + quote(value, safe=""),
                    customer=customer_by_so.get(value.casefold(), ""),
                    ship_date=ship_date_by_so.get(value.casefold(), ""),
                )
        if "Name" in so.columns:
            for value in so["Name"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist():
                add("Customer", value, "/?customer=" + quote(value, safe=""))
        if "Item" in so.columns:
            for value in so["Item"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist():
                add(
                    "Item",
                    value,
                    "/quotation_lookup?item=" + quote(value, safe=""),
                    on_hand=on_hand_by_item.get(value.casefold(), ""),
                )

    if inventory is not None and not inventory.empty and "Part_Number" in inventory.columns:
        for value in inventory["Part_Number"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist():
            add(
                "Item",
                value,
                "/quotation_lookup?item=" + quote(value, safe=""),
                on_hand=on_hand_by_item.get(value.casefold(), ""),
            )

    return entries

def _load_from_db(force: bool = False):
    global SO_INV, INVENTORY_STATUS, SAP, OPEN_PO, LEDGER, ITEM_ATP, _LAST_LOAD_ERR, _LAST_LOADED_AT
    global ITEM_SUGGEST_CACHE, GLOBAL_SEARCH_INDEX
    global SO_LOOKUP_BASE, WAITING_ITEMS_BY_QB, LEDGER_ITEM_INDEX
    global PDF_DB_SEARCH_CACHE, INDEX_VIEW_CACHE, QUOTATION_VIEW_CACHE, QUOTE_ITEM_SUGGEST_ROWS, READY_ASSIGN_CACHE
    try:
        if (
            force
            or SO_INV is None
            or SAP is None
            or OPEN_PO is None
            or LEDGER is None
            or ITEM_ATP is None
        ):
            so = _read_table("public", "wo_structured")
            inventory = _read_table("public", "inventory_status")
            sap = _read_table("public", "NT Shipping Schedule")
            open_po = _read_table("public", "Open_Purchase_Orders")
            ledger = _read_table("public", "ledger_analytics")
            for c in ("Ship Date", "Order Date"):
                _safe_date_col(so, c)
                _safe_date_col(sap, c)
            for col in open_po.columns:
                if "date" in col.lower():
                    _safe_date_col(open_po, col)
            if "Date" in ledger.columns:
                _safe_date_col(ledger, "Date")

            SO_INV, INVENTORY_STATUS, SAP, OPEN_PO = so, inventory, sap, open_po
            LEDGER = ledger
            ITEM_ATP = build_atp_view(ledger)
            SO_LOOKUP_BASE, WAITING_ITEMS_BY_QB, LEDGER_ITEM_INDEX = _build_runtime_indexes(so, ledger)
            suggest_items: list[str] = []
            if "Item" in so.columns:
                suggest_items.extend(
                    so["Item"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].tolist()
                )
            if "Part_Number" in inventory.columns:
                suggest_items.extend(
                    inventory["Part_Number"].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].tolist()
                )
            ITEM_SUGGEST_CACHE = sorted(set(suggest_items))
            GLOBAL_SEARCH_INDEX = _build_global_search_index(so, inventory)
            QUOTE_ITEM_SUGGEST_ROWS = _build_quote_item_summaries(inventory, ledger)
            PDF_DB_SEARCH_CACHE = {}
            INDEX_VIEW_CACHE = {}
            QUOTATION_VIEW_CACHE = {}
            READY_ASSIGN_CACHE = None
            _LAST_LOAD_ERR = None
            _LAST_LOADED_AT = datetime.now()
    except Exception as e:
        SO_INV = None
        INVENTORY_STATUS = None
        SAP = None
        OPEN_PO = None
        LEDGER = None
        ITEM_ATP = None
        SO_LOOKUP_BASE = None
        WAITING_ITEMS_BY_QB = {}
        LEDGER_ITEM_INDEX = {}
        ITEM_SUGGEST_CACHE = []
        GLOBAL_SEARCH_INDEX = []
        QUOTE_ITEM_SUGGEST_ROWS = []
        PDF_DB_SEARCH_CACHE = {}
        INDEX_VIEW_CACHE = {}
        QUOTATION_VIEW_CACHE = {}
        READY_ASSIGN_CACHE = None
        _LAST_LOAD_ERR = f"DB load error: {e}"

def _ensure_loaded():
    if (
        SO_INV is None
        or INVENTORY_STATUS is None
        or SAP is None
        or OPEN_PO is None
        or LEDGER is None
        or ITEM_ATP is None
    ):
        _load_from_db(force=True)
    # Load PDF map on demand as well
    _load_pdf_map()


def _ready_to_assign_rows() -> list[dict]:
    global READY_ASSIGN_CACHE
    if READY_ASSIGN_CACHE is not None:
        return READY_ASSIGN_CACHE

    if SO_INV is None or SO_INV.empty or ITEM_ATP is None or ITEM_ATP.empty:
        READY_ASSIGN_CACHE = []
        return READY_ASSIGN_CACHE

    placeholder_date = PLACEHOLDER_DATE
    today = pd.Timestamp.today().normalize()

    so = SO_INV.copy()
    for c in ["QB Num", "Item", "Qty(-)", "Ship Date", "Name", "P. O. #", "Order Date", "Component_Status"]:
        if c not in so.columns:
            so[c] = pd.NA
    so["Ship Date"] = pd.to_datetime(so["Ship Date"], errors="coerce")
    so["Order Date"] = pd.to_datetime(so["Order Date"], errors="coerce")
    so["QB Num"] = so["QB Num"].astype(str).str.strip()
    so["Item"] = so["Item"].astype(str).str.strip()
    so["Qty(-)"] = pd.to_numeric(so["Qty(-)"], errors="coerce").fillna(0.0)
    pending = so.loc[
        so["Ship Date"].eq(placeholder_date) & so["QB Num"].ne("") & so["Item"].ne("")
    ].copy()
    if pending.empty:
        READY_ASSIGN_CACHE = []
        return READY_ASSIGN_CACHE

    rows: list[dict] = []
    for qb_num, grp in pending.groupby("QB Num", sort=True):
        demands = (
            grp.loc[grp["Qty(-)"] > 0, ["Item", "Qty(-)"]]
            .groupby("Item", as_index=False)["Qty(-)"]
            .sum()
        )
        if demands.empty:
            continue
        demand_map = {str(r["Item"]): float(r["Qty(-)"]) for _, r in demands.iterrows()}
        ready_dt = earliest_atp_for_items_strict(ITEM_ATP, demand_map, from_date=today, allow_zero=True)
        if ready_dt is None or ready_dt >= placeholder_date:
            continue

        first = grp.iloc[0]
        waiting_items = sorted(
            set(
                grp.loc[grp["Component_Status"].isin(["Waiting", "Shortage"]), "Item"]
                .dropna()
                .astype(str)
                .str.strip()
                .loc[lambda s: s.ne("")]
                .tolist()
            )
        )
        rows.append(
            {
                "qb_num": str(qb_num),
                "customer": str(first.get("Name") or ""),
                "po_num": str(first.get("P. O. #") or ""),
                "order_date": first["Order Date"].strftime("%Y-%m-%d") if pd.notna(first.get("Order Date")) else "",
                "current_ship_date": placeholder_date.strftime("%Y-%m-%d"),
                "ready_date": ready_dt.strftime("%Y-%m-%d"),
                "item_count": int(len(demand_map)),
                "waiting_items": ", ".join(waiting_items),
            }
        )

    rows.sort(key=lambda r: (r["ready_date"], r["qb_num"]))
    READY_ASSIGN_CACHE = rows
    return READY_ASSIGN_CACHE


def _push_recent_home_search(*, so_input: str = "", customer_input: str = "") -> None:
    global RECENT_HOME_SEARCHES

    so_value = str(so_input or "").strip()
    customer_value = str(customer_input or "").strip()
    if not so_value and not customer_value:
        return

    if so_value:
        entry = {"kind": "SO", "label": so_value, "href": f"/?so={so_value}"}
    else:
        entry = {"kind": "Customer", "label": customer_value, "href": f"/?customer={customer_value}"}

    dedup_key = (entry["kind"], entry["label"].strip().upper())
    RECENT_HOME_SEARCHES = [
        item for item in RECENT_HOME_SEARCHES
        if (item.get("kind", ""), str(item.get("label", "")).strip().upper()) != dedup_key
    ]
    RECENT_HOME_SEARCHES.insert(0, entry)
    RECENT_HOME_SEARCHES = RECENT_HOME_SEARCHES[:5]


def _dashboard_lt_unassigned_count() -> int:
    if SO_INV is None or SO_INV.empty or "Ship Date" not in SO_INV.columns or "QB Num" not in SO_INV.columns:
        return 0
    so = SO_INV.copy()
    so["Ship Date"] = pd.to_datetime(so["Ship Date"], errors="coerce")
    so["QB Num"] = so["QB Num"].astype(str).str.strip()
    unassigned_mask = (
        (so["Ship Date"].dt.month.eq(7) & so["Ship Date"].dt.day.eq(4))
        | (so["Ship Date"].dt.month.eq(12) & so["Ship Date"].dt.day.eq(31))
    )
    return int(so.loc[unassigned_mask & so["QB Num"].ne(""), "QB Num"].nunique())


def _dashboard_top_shortage_items(limit: int = 5) -> list[dict[str, object]]:
    if SO_INV is None or SO_INV.empty or "Item" not in SO_INV.columns or "QB Num" not in SO_INV.columns:
        return []

    so = SO_INV.copy()
    for col in ("Component_Status", "Qty(-)", "On Hand"):
        if col not in so.columns:
            so[col] = 0 if col != "Component_Status" else ""

    so["Item"] = so["Item"].astype(str).str.strip()
    so["QB Num"] = so["QB Num"].astype(str).str.strip()
    so["Qty(-)"] = pd.to_numeric(so["Qty(-)"], errors="coerce").fillna(0.0)
    so["On Hand"] = pd.to_numeric(so["On Hand"], errors="coerce").fillna(0.0)

    shortage = so.loc[
        so["Component_Status"].isin(["Waiting", "Shortage"])
        & so["Item"].ne("")
        & so["QB Num"].ne("")
    ].copy()
    if shortage.empty:
        return []

    grouped = (
        shortage.groupby("Item", as_index=False)
        .agg(
            blocked_so_count=("QB Num", "nunique"),
            open_so_qty=("Qty(-)", "sum"),
            on_hand=("On Hand", "max"),
        )
        .sort_values(["blocked_so_count", "open_so_qty", "Item"], ascending=[False, False, True])
        .head(limit)
    )

    out: list[dict[str, object]] = []
    for _, row in grouped.iterrows():
        out.append(
            {
                "item": str(row["Item"]),
                "blocked_so_count": int(row["blocked_so_count"]),
                "open_so_qty": _coerce_total(row["open_so_qty"]) or 0,
                "on_hand": _coerce_total(row["on_hand"]) or 0,
            }
        )
    return out


def _dashboard_alerts() -> list[dict[str, str]]:
    alerts: list[dict[str, str]] = []


    if OPEN_PO is not None and not OPEN_PO.empty:
        pod = OPEN_PO.copy()
        pod_no_col = "POD#" if "POD#" in pod.columns else ("QB Num" if "QB Num" in pod.columns else None)
        ship_col = "Ship Date" if "Ship Date" in pod.columns else None
        if pod_no_col and ship_col:
            pod[pod_no_col] = pod[pod_no_col].fillna("").astype(str).str.strip()
            pod[ship_col] = pd.to_datetime(pod[ship_col], errors="coerce")
            missing_ship_count = pod.loc[pod[pod_no_col].ne("") & pod[ship_col].isna(), pod_no_col].nunique()
            if missing_ship_count:
                alerts.append({"label": f"{int(missing_ship_count)} PODs are missing ship dates.", "href": ""})

    if LEDGER is not None and not LEDGER.empty and {"Date", "Projected_NAV", "Item"}.issubset(LEDGER.columns):
        led = LEDGER.copy()
        led["Date"] = pd.to_datetime(led["Date"], errors="coerce")
        led["Projected_NAV"] = pd.to_numeric(led["Projected_NAV"], errors="coerce")
        placeholder_date = PLACEHOLDER_DATE
        neg_item_count = (
            led.loc[
                led["Date"].notna()
                & led["Date"].ne(placeholder_date)
                & led["Projected_NAV"].lt(0),
                "Item",
            ]
            .dropna().astype(str).str.strip().loc[lambda s: s.ne("")].nunique()
        )
        if neg_item_count:
            alerts.append({"label": f"{int(neg_item_count)} items will go negative in the future", "href": "/dashboard/negative_inventory"})

    if _LAST_LOADED_AT is not None:
        age_minutes = max(0, int((datetime.now() - _LAST_LOADED_AT).total_seconds() // 60))
        if age_minutes >= 60:
            alerts.append({"label": f"Homepage data is {age_minutes} minutes old.", "href": "/?reload=1"})

    if not alerts:
        alerts.append({"label": "No active system alerts.", "href": ""})
    return alerts[:4]

def _negative_inventory_detail_rows(limit: int | None = None) -> tuple[list[str], list[dict[str, object]]]:
    columns = ["Item", "Date", "Projected Qty"]
    if LEDGER is None or LEDGER.empty or not {"Date", "Projected_NAV", "Item"}.issubset(LEDGER.columns):
        return columns, []

    led = LEDGER.copy()
    led["Date"] = pd.to_datetime(led["Date"], errors="coerce")
    led["Projected_NAV"] = pd.to_numeric(led["Projected_NAV"], errors="coerce")
    placeholder_date = PLACEHOLDER_DATE
    neg = led.loc[
        led["Date"].notna()
        & led["Date"].ne(placeholder_date)
        & led["Projected_NAV"].lt(0)
        & led["Item"].notna(),
        ["Item", "Date", "Projected_NAV"],
    ].copy()
    if neg.empty:
        return columns, []

    neg = neg.sort_values(["Date", "Item", "Projected_NAV"], kind="mergesort")
    if limit is not None:
        neg = neg.head(limit)

    rows = [
        {
            "Item": str(row["Item"]),
            "Date": row["Date"].strftime("%Y-%m-%d") if pd.notna(row["Date"]) else "",
            "Projected Qty": _format_intish(row["Projected_NAV"]),
        }
        for _, row in neg.iterrows()
    ]
    return columns, rows

def lookup_on_po_by_item(item: str) -> int | None:
    df = SO_INV[SO_INV["Item"] == item]
    if "On PO" not in df.columns:
        return None
    s = pd.to_numeric(df["On PO"], errors="coerce").dropna()
    return int(s.iloc[0]) if not s.empty else None

def lookup_on_sales_by_item(item: str) -> int | float | None:
    df = SO_INV[SO_INV["Item"] == item]
    col_name = None
    for candidate in ("On Sales Order", "On Sales", "On SO"):
        if candidate in df.columns:
            col_name = candidate
            break
    if not col_name:
        return None
    s = pd.to_numeric(df[col_name], errors="coerce").dropna()
    if s.empty:
        return None
    first = s.iloc[0]
    return int(first) if s.eq(first).all() else int(s.sum())

def _coerce_total(val):
    if pd.isna(val):
        return None
    as_float = float(val)
    return int(as_float) if as_float.is_integer() else as_float

def _format_intish(val: object) -> str:
    if val is None or val == "" or pd.isna(val):
        return ""
    try:
        fval = float(val)
    except Exception:
        return str(val)
    return str(int(fval)) if fval.is_integer() else str(fval)

def _aggregate_metric(series: pd.Series) -> int | float | None:
    numeric = pd.to_numeric(series, errors="coerce").dropna()
    if numeric.empty:
        return None
    first = numeric.iloc[0]
    if numeric.eq(first).all():
        return _coerce_total(first)
    total = numeric.sum()
    return _coerce_total(total)


def _resolve_ledger_item_key(item: str) -> str:
    raw = str(item or "").strip().upper()
    if not raw:
        return ""
    if raw in LEDGER_ITEM_INDEX:
        return raw
    for key in LEDGER_ITEM_INDEX.keys():
        if str(key).upper() == raw:
            return str(key)
    return raw


def _lookup_earliest_atp_date(item: str, qty: float = 1.0) -> datetime | None:
    """Look up ATP in the view rebuilt from the ledger during data loading."""
    if ITEM_ATP is None or ITEM_ATP.empty:
        return None
    atp_dt = earliest_atp_strict(
        ITEM_ATP, item, qty,
        from_date=pd.Timestamp.today().normalize(), allow_zero=True,
    )
    return None if atp_dt is None else atp_dt.to_pydatetime()


def _find_pdf_url_for_so(so_num: str, po_num: str | None = None) -> str | None:
    """
    Best-effort PDF link lookup for a given SO/QB number.
    Mirrors the logic used on the main index page.
    """
    so_num = (so_num or "").strip()
    if not so_num:
        return None

    search_keys = [so_num]
    so_upper = so_num.upper()
    if so_upper.startswith("SO-"):
        search_keys.append(so_num[3:])
    if po_num:
        search_keys.append(str(po_num))

    pdf_record = None
    for key in search_keys:
        recs = _pdf_db_search_by_filename(key, limit=1)
        if recs:
            pdf_record = recs[0]
            break

    if pdf_record:
        view_url = _pdf_view_url_for_path(pdf_record.get("file_path"))
        if view_url:
            return view_url
        return f"/pdfid/{pdf_record['id']}"

    keys_to_try = [so_num, so_upper.replace("SO-", ""), so_upper.replace("SO", "").strip("- ")]
    pdf_info = None
    for k in keys_to_try:
        pdf_info = PDF_MAP.get(k.upper())
        if pdf_info:
            break
    if pdf_info:
        return f"/pdf/{(pdf_info['file_name'][:-4])}"
    return None


def _wo_status_by_qb_num() -> dict[str, str]:
    if SO_INV is None or SO_INV.empty or "QB Num" not in SO_INV.columns or "Picked" not in SO_INV.columns:
        return {}

    so = SO_INV[["QB Num", "Picked"]].copy()
    so["QB Num"] = so["QB Num"].fillna("").astype(str).str.strip()
    so["Picked"] = so["Picked"].fillna("").astype(str).str.strip()
    so = so.loc[so["QB Num"].ne("")].copy()
    if so.empty:
        return {}

    def _status_for_group(series: pd.Series) -> str:
        values = {str(v).strip().lower() for v in series if str(v).strip()}
        return "Picked" if ("picked" in values or "partial" in values) else "NA"

    return so.groupby("QB Num")["Picked"].agg(_status_for_group).to_dict()

def _so_table_for_item(item: str) -> tuple[list[str], list[dict], dict[str, int | float | None]]:
    need_cols = ["Name", "QB Num", "Item", "Qty(-)", "On Hand - WIP", "Ship Date", "Picked"]
    g = SO_INV[SO_INV["Item"] == item].copy()
    for c in need_cols:
        if c not in g.columns:
            g[c] = ""
    # Fallback for WIP column if missing in data
    if "On Hand - WIP" not in SO_INV.columns and "In Stock(Inventory)" in SO_INV.columns:
        g["On Hand - WIP"] = SO_INV.loc[g.index, "In Stock(Inventory)"]
    if "Ship Date" in g.columns:
        ship_dates = pd.to_datetime(g["Ship Date"], errors="coerce")
        g = (
            g.assign(_ship_date_sort=ship_dates)
            .sort_values("_ship_date_sort", na_position="last")
            .drop(columns="_ship_date_sort")
        )
        g["Ship Date"] = _to_date_str(g["Ship Date"])
    rows = g[need_cols].fillna("").astype(str).to_dict(orient="records") if not g.empty else []
    totals = {"on_sales_order": None, "on_po": None}
    if not g.empty:
        if "On Sales Order" in g.columns:
            totals["on_sales_order"] = _aggregate_metric(g["On Sales Order"])
        if "On PO" in g.columns:
            totals["on_po"] = _aggregate_metric(g["On PO"])
    return need_cols, rows, totals

def _compute_on_hand_metrics(df: pd.DataFrame) -> tuple[int | float | None, int | float | None]:
    if df is None or df.empty:
        return None, None
    on_hand = _aggregate_metric(df.get("On Hand", pd.Series(dtype=float)))
    # Fall back if "On Hand - WIP" is missing
    col_wip = "On Hand - WIP" if "On Hand - WIP" in df.columns else ("In Stock(Inventory)" if "In Stock(Inventory)" in df.columns else None)
    on_hand_wip = _aggregate_metric(df.get(col_wip, pd.Series(dtype=float))) if col_wip else None
    return on_hand, on_hand_wip


def _item_lookup_values(item: str) -> set[str]:
    raw = str(item or "").strip()
    values = {raw, raw.upper()}
    try:
        norm = str(normalize_item(raw)).strip()
        values.update({norm, norm.upper()})
    except Exception:
        pass
    return {v for v in values if v}


def _first_existing_column(df: pd.DataFrame, candidates: tuple[str, ...]) -> str | None:
    lookup = {str(col).lower(): col for col in df.columns}
    for candidate in candidates:
        found = lookup.get(candidate.lower())
        if found is not None:
            return found
    return None



def _recent_receiving_summary_for_item(item: str, days: int = 7) -> tuple[list[str], list[dict]]:
    global RECEIVING_LOG
    if not item:
        return [], []
    if RECEIVING_LOG is None:
        RECEIVING_LOG = pd.DataFrame()
        for receiving_table in ("receving_log", "receiving_log", "receving-log", "receiving-log"):
            try:
                RECEIVING_LOG = _read_table("public", receiving_table)
                break
            except Exception:
                continue
    if RECEIVING_LOG.empty:
        return [], []

    df = RECEIVING_LOG.copy()
    item_col = _first_existing_column(
        df,
        ("part_number", "Part_Number", "Part Number", "Item", "item", "part_name", "Part Name"),
    )
    date_col = _first_existing_column(df, ("entry_date", "Entry Date", "date", "Date", "received_date", "Received Date"))
    qty_col = _first_existing_column(df, ("quantity", "Quantity", "qty", "Qty", "received_qty", "Received Qty"))
    inv_col = _first_existing_column(df, ("invoice_number", "Invoice Number", "Invoice#", "Inv#", "Inv #"))
    pod_col = _first_existing_column(df, ("pod_number", "POD Number", "POD#", "POD #", "pod"))
    ref_col = _first_existing_column(df, ("Reference", "reference", "Ref", "ref"))
    if item_col is None or date_col is None or qty_col is None:
        return [], []

    wanted = {v.upper() for v in _item_lookup_values(item)}
    df[item_col] = df[item_col].fillna("").astype(str).str.strip()
    df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
    df[qty_col] = pd.to_numeric(df[qty_col], errors="coerce").fillna(0.0)
    today = pd.Timestamp.today().normalize()
    start = today - pd.Timedelta(days=days)
    mask = df[item_col].str.upper().isin(wanted) & df[date_col].notna() & (df[date_col] >= start) & (df[date_col] < today + pd.Timedelta(days=1))
    df = df.loc[mask].copy()
    if df.empty:
        return ["Date", "Inv#", "POD#", "Reference", "Qty"], []

    df["Date"] = df[date_col].dt.strftime("%Y-%m-%d")
    df["Inv#"] = df[inv_col].fillna("").astype(str).str.strip() if inv_col else ""
    df["POD#"] = df[pod_col].fillna("").astype(str).str.strip() if pod_col else ""
    df["Reference"] = df[ref_col].fillna("").astype(str).str.strip() if ref_col else ""
    grouped = (
        df.groupby(["Date", "Inv#", "POD#", "Reference"], dropna=False, as_index=False)[qty_col]
        .sum()
        .rename(columns={qty_col: "Qty"})
    )
    grouped["Qty"] = grouped["Qty"].apply(_format_intish)
    columns = ["Date", "Inv#", "POD#", "Reference", "Qty"]
    rows = grouped.sort_values(["Date", "Inv#", "POD#", "Reference"], ascending=[False, True, True, True], kind="mergesort").fillna("").astype(str).to_dict(orient="records")
    return columns, rows

def _po_table_for_item(item: str) -> tuple[list[str], list[dict]]:
    if "Item" not in SAP.columns:
        raise ValueError("SAP table missing 'Item' column.")
    item_lower = item.lower()
    item_upper = item.upper()
    sap_item_series = SAP["Item"].astype(str)
    mask = sap_item_series.str.lower() == item_lower
    allow_desc_lookup = not item_upper.startswith(("N", "SEMIL", "POC"))
    if allow_desc_lookup and "Description" in SAP.columns:
        desc_mask = SAP["Description"].astype(str).str.lower().str.contains(item_lower, na=False)
        mask |= desc_mask
    g = SAP[mask].copy()
    for dc in ("Ship Date", "Order Date", "ETA"):
        if dc in g.columns:
            g[dc] = _to_date_str(g[dc])
    cols = list(g.columns) if not g.empty else list(SAP.columns)
    g = g.fillna("").astype(str)
    rows = g[cols].to_dict(orient="records") if not g.empty else []
    return cols, rows

def _open_po_table_for_item(item: str) -> tuple[list[str], list[dict]]:
    if OPEN_PO is None or OPEN_PO.empty:
        return [], []

    item_lower = item.lower()
    item_upper = item.upper()

    df = OPEN_PO
    item_col = next((c for c in df.columns if c.lower() == "item"), None)
    desc_col = next((c for c in df.columns if c.lower() == "description"), None)

    hide_cols = {"deliv date", "ship date"}
    if item_col is None and desc_col is None:
        cols = [c for c in df.columns if c.lower() not in hide_cols]
        return cols, []

    mask = pd.Series(False, index=df.index)
    if item_col:
        mask |= df[item_col].astype(str).str.lower() == item_lower

    allow_desc_lookup = not item_upper.startswith(("N", "SEMIL", "POC"))
    if allow_desc_lookup and desc_col:
        mask |= df[desc_col].astype(str).str.lower().str.contains(item_lower, na=False)

    result = df.loc[mask].copy()
    if result.empty:
        cols = [c for c in df.columns if c.lower() not in hide_cols]
        return cols, []

    for col in result.columns:
        if "date" in col.lower():
            _safe_date_col(result, col)

    for col in list(result.columns):
        if col.lower() in hide_cols:
            result.drop(columns=[col], inplace=True)

    result = result.fillna("").astype(str)
    return list(result.columns), result.to_dict(orient="records")


# initial load
_load_from_db(force=True)

# =========================
# Routes
# =========================
@app.route("/", methods=["GET", "POST"])
def index():
    if request.args.get("reload") == "1":
        _load_from_db(force=True)
        _load_pdf_map(force=True)
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    lt_unassigned_count = _dashboard_lt_unassigned_count()
    alerts = _dashboard_alerts()
    recent_searches = list(RECENT_HOME_SEARCHES)

    # ---- read inputs (work with GET or POST) ----
    so_input = (request.values.get("so") or "").strip()
    customer_input = (request.values.get("customer") or "").strip()
    cache_key = (so_input, customer_input)
    if cache_key in INDEX_VIEW_CACHE:
        cached = INDEX_VIEW_CACHE[cache_key]
        return render_template_string(
            INDEX_TPL,
            so_num=so_input,
            customer_val=customer_input,
            customer_query=cached["customer_query"],
            customer_options=cached["customer_options"],
            rows=cached["rows"],
            count=cached["count"],
            loaded_at=_LAST_LOADED_AT.strftime("%Y-%m-%d %H:%M:%S") if _LAST_LOADED_AT else "â€”",
            order_summary=cached["order_summary"],
            lt_unassigned_count=lt_unassigned_count,
            recent_searches=recent_searches,
            alerts=alerts,
            headers=cached["table_headers"],
            header_labels=TABLE_HEADER_LABELS,
            numeric_cols=[
                "Qty(-)","Available","Available + Pre-installed PO","On Hand",
                "On Sales Order","On PO","Assigned Q'ty","On Hand - WIP",
                "Available + On PO","Sales/Week","Recommended Restock Qty"
            ],
        )

    # Flexible SO handling: allow "20251368", "SO20251368", or "so-20251368"
    so_num = ""
    so_upper = so_input.upper()
    if so_upper:
        if so_upper.startswith("SO-"):
            so_num = so_upper
        elif so_upper.startswith("SO") and so_upper[2:].replace("-", "").isdigit():
            so_num = f"SO-{so_upper[2:].lstrip('-')}"
        elif so_upper.replace("-", "").isdigit():
            so_num = f"SO-{so_upper.replace('-', '')}"

    rows, count = None, 0
    order_summary = None
    table_headers = None
    customer_options = None
    customer_query = None
    so_base = SO_LOOKUP_BASE if SO_LOOKUP_BASE is not None else SO_INV
    if so_input:
        rows_df = pd.DataFrame()

        if so_num:
            mask = so_base["__qb_upper"] == so_num if "__qb_upper" in so_base.columns else so_base["QB Num"].astype(str).str.upper() == so_num
            rows_df = SO_INV.loc[mask].copy()

            # Fallback: tolerant key match for formatting differences.
            if rows_df.empty and "QB Num" in SO_INV.columns:
                target_key = _normalize_so_key(so_num)
                qb_norm = SO_INV["QB Num"].astype(str).map(_normalize_so_key)
                rows_df = SO_INV.loc[qb_norm.eq(target_key)].copy()

        if (rows_df is None or rows_df.empty) and "Name" in SO_INV.columns:
            name_col = so_base["__name_text"] if "__name_text" in so_base.columns else so_base["Name"].astype(str)
            name_mask = name_col.str.contains(so_input, case=False, na=False)
            rows_df = SO_INV.loc[name_mask].copy()

        count = len(rows_df)

        if "On Hand - WIP" not in rows_df.columns and "In Stock(Inventory)" in rows_df.columns:
            rows_df["On Hand - WIP"] = rows_df["In Stock(Inventory)"]

        required_headers = [
            "Order Date","Name","P. O. #","QB Num","Item","Component_Status","Qty(-)","Available",
            "Available + On PO","Available + Pre-installed PO","On Hand","On Sales Order","On PO",
            "Assigned Q'ty","On Hand - WIP","Sales/Week",
            "Recommended Restock Qty","Ship Date"
        ]
        for h in required_headers:
            if h not in rows_df.columns: rows_df[h] = ""

        qb_for_pdf = None
        if "QB Num" in rows_df.columns:
            qb_vals = rows_df["QB Num"].dropna().astype(str).unique().tolist()
            if len(qb_vals) == 1:
                qb_for_pdf = qb_vals[0]

        summary_cols = ["Order Date", "Name", "P. O. #", "QB Num", "Ship Date"]
        summary_fields = []
        for col in summary_cols:
            col_vals = rows_df[col].dropna().astype(str) if col in rows_df.columns else pd.Series(dtype=str)
            summary_fields.append({
                "label": col,
                "value": col_vals.iloc[0] if not col_vals.empty else "",
            })
        order_summary = {
            "qb_num": qb_for_pdf or so_input,
            "row_count": count,
            "fields": summary_fields,
        }

        # Attach PDF link by DB search first (ILIKE on filename); fallback to filesystem map
        po_num = ""
        if "P. O. #" in rows_df.columns:
            ser = rows_df["P. O. #"].dropna().astype(str)
            po_num = ser.iloc[0] if not ser.empty else ""

        search_keys = []
        if qb_for_pdf:
            search_keys.append(qb_for_pdf)
            qb_upper = qb_for_pdf.upper()
            if qb_upper.startswith("SO-"):
                search_keys.append(qb_upper[3:])  # numeric only
            if po_num:
                search_keys.append(str(po_num))

        pdf_record = None
        if search_keys:
            for key in search_keys:
                recs = _pdf_db_search_by_filename(key, limit=1)
                if recs:
                    pdf_record = recs[0]
                    break

        if pdf_record:
            order_summary["pdf_url"] = f"/pdfid/{pdf_record['id']}"
            order_summary["pdf_name"] = pdf_record.get("file_name")
        else:
            # Fallback to filesystem map (exact filename stems)
            keys_to_try = []
            if qb_for_pdf:
                keys_to_try = [qb_for_pdf, qb_for_pdf.replace("SO-", ""), qb_for_pdf.replace("SO", "").strip("- ")]
            pdf_info = None
            for k in keys_to_try:
                pdf_info = PDF_MAP.get(k.upper())
                if pdf_info:
                    break
            if pdf_info:
                order_summary["pdf_url"] = f"/pdf/{(pdf_info['file_name'][:-4])}"
                order_summary["pdf_name"] = pdf_info["file_name"]

        for c in ("Ship Date", "Order Date"):
            if c in rows_df.columns: rows_df[c] = _to_date_str(rows_df[c])
        table_headers = [h for h in required_headers if h not in ("Order Date","Name","P. O. #","QB Num","Ship Date")]
        table_df = rows_df[table_headers].copy()
        rows = table_df.fillna("").astype(str).to_dict(orient="records")
    elif customer_input:
        customer_query = customer_input
        customer_options = []
        if "Name" in SO_INV.columns:
            name_col = so_base["__name_text"] if "__name_text" in so_base.columns else so_base["Name"].astype(str)
            name_mask = name_col.str.contains(customer_input, case=False, na=False)
            cust_df = SO_INV.loc[name_mask].copy()
            if not cust_df.empty:
                if "QB Num" not in cust_df.columns:
                    cust_df["QB Num"] = ""
                for c in ("Ship Date", "Order Date"):
                    if c in cust_df.columns:
                        cust_df[c] = _to_date_str(cust_df[c])
                for qb_num, grp in cust_df.groupby("QB Num"):
                    qb_str = str(qb_num).strip()
                    if not qb_str:
                        continue
                    first = grp.iloc[0]
                    customer_options.append({
                        "qb_num": qb_str,
                        "name": first.get("Name", ""),
                        "ship_date": first.get("Ship Date", ""),
                        "order_date": first.get("Order Date", ""),
                    })
                customer_options.sort(key=lambda x: x.get("qb_num", ""))

    if so_input or customer_input:
        _push_recent_home_search(so_input=so_input, customer_input=customer_input)
        recent_searches = list(RECENT_HOME_SEARCHES)

    cache_payload = {
        "rows": rows,
        "count": count,
        "order_summary": order_summary,
        "table_headers": table_headers,
        "customer_options": customer_options,
        "customer_query": customer_query,
    }
    if len(INDEX_VIEW_CACHE) >= 256:
        INDEX_VIEW_CACHE.clear()
    INDEX_VIEW_CACHE[cache_key] = cache_payload

    return render_template_string(
        INDEX_TPL,
        so_num=so_input,           # show original entry
        customer_val=customer_input,
        customer_query=customer_query,
        customer_options=customer_options,
        rows=rows,
        count=count,
        loaded_at=_LAST_LOADED_AT.strftime("%Y-%m-%d %H:%M:%S") if _LAST_LOADED_AT else "â€”",
        order_summary=order_summary,
        lt_unassigned_count=lt_unassigned_count,
        recent_searches=recent_searches,
        alerts=alerts,
        headers=table_headers,
        header_labels=TABLE_HEADER_LABELS,
        numeric_cols=[
            "Qty(-)","Available","Available + Pre-installed PO","On Hand",
            "On Sales Order","On PO","Assigned Q'ty","On Hand - WIP",
            "Available + On PO","Sales/Week","Recommended Restock Qty"
        ],
    )

@app.route("/dashboard/negative_inventory")
def dashboard_negative_inventory():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    columns, rows = _negative_inventory_detail_rows()
    return render_template_string(
        """
<!doctype html>
<html>
<head>
  <link rel="icon" href="/static/robot.svg?v=1" type="image/svg+xml">
  <meta charset="utf-8">
  <title>Negative Inventory Details</title>
  <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
  <style>
    body{ background:#f7fafc; color:#0f172a; padding:28px; }
    .card-lite{ border-radius:14px; box-shadow:0 10px 22px rgba(0,0,0,.06); }
    .table-responsive{ max-height:78vh; overflow:auto; }
    .table thead th{ position:sticky; top:0; z-index:2; background:#f8fafc; white-space:nowrap; }
    .num{ text-align:right; font-variant-numeric:tabular-nums; }
  </style>
</head>
<body>
  <div class="d-flex justify-content-between align-items-center mb-3">
    <div>
      <div class="h3 m-0">Negative Inventory Details</div>
      <div class="text-muted small">Negative projected quantity rows with real scheduled dates.</div>
    </div>
    <a class="btn btn-sm btn-outline-secondary" href="/">Home</a>
  </div>
  <div class="card-lite bg-white p-3">
    <div class="d-flex justify-content-between align-items-center mb-2">
      <div class="fw-bold">{{ rows|length }} row(s)</div>
      <div class="text-muted small">Source: public.ledger_analytics</div>
    </div>
    <div class="table-responsive">
      <table class="table table-sm table-bordered table-hover align-middle mb-0">
        <thead class="table-light text-uppercase small text-muted">
          <tr>{% for c in columns %}<th>{{ c }}</th>{% endfor %}</tr>
        </thead>
        <tbody>
          {% if rows %}
            {% for row in rows %}
              <tr>
                {% for c in columns %}
                  <td class="{{ 'num' if c == 'Projected Qty' else '' }}">{{ row[c] }}</td>
                {% endfor %}
              </tr>
            {% endfor %}
          {% else %}
            <tr><td colspan="{{ columns|length }}" class="text-center text-muted py-4">No negative projected quantity rows found.</td></tr>
          {% endif %}
        </tbody>
      </table>
    </div>
  </div>
</body>
</html>
        """,
        columns=columns,
        rows=rows,
    )
@app.route("/api/reload", methods=["POST"])
def api_reload():
    _load_from_db(force=True)
    _load_pdf_map(force=True)
    if _LAST_LOAD_ERR:
        return jsonify({"ok": False, "error": _LAST_LOAD_ERR}), 500
    return jsonify({"ok": True, "loaded_at": _LAST_LOADED_AT.isoformat()})


@app.route("/api/debug/db_state")
def api_debug_db_state():
    _ensure_loaded()
    so = (request.args.get("so") or "").strip()
    so_count = None
    if so and SO_INV is not None and not SO_INV.empty and "QB Num" in SO_INV.columns:
        so_count = int(SO_INV["QB Num"].astype(str).str.upper().eq(so.upper()).sum())
    return jsonify(
        {
            "ok": _LAST_LOAD_ERR is None,
            "loaded_at": _LAST_LOADED_AT.isoformat() if _LAST_LOADED_AT else None,
            "database_dsn": str(engine.url),
            "pdf_folder": PDF_FOLDER,
            "so_inv_rows": int(len(SO_INV)) if SO_INV is not None else None,
            "wo_structured_match_count": so_count,
            "last_load_error": _LAST_LOAD_ERR,
        }
    )

@app.route("/pdf/<order_id>")
def serve_pdf(order_id: str):
    """Serve a PDF by order id (stem of filename). Only serves files under PDF_FOLDER.
    """
    _load_pdf_map()
    if not PDF_FOLDER:
        abort(404)
    info = PDF_MAP.get(order_id.upper())
    if not info:
        # try variants: with SO- prefix or stripped
        variants = [order_id, f"SO-{order_id}", order_id.replace("SO-", ""), order_id.replace("SO", "").strip("- ")]
        for v in variants:
            info = PDF_MAP.get(v.upper())
            if info:
                break
    if not info:
        abort(404)
    path = info["file_path"]
    if not os.path.isfile(path):
        abort(404)
    # Send file directly; let browser handle PDF
    return send_file(path, mimetype="application/pdf", as_attachment=False, download_name=info["file_name"])

@app.route("/favicon.ico")
def favicon():
    """Use the same robot icon for browsers requesting the conventional URL."""
    return redirect(url_for("static", filename="robot.svg", v=1))

@app.route("/pdfid/<int:pdf_id>")
def serve_pdf_by_id(pdf_id: int):
    """Serve a PDF using a record in pdf_file_log by id."""
    rec = _pdf_db_get_by_id(pdf_id)
    if not rec:
        abort(404)
    path = rec.get("file_path")
    name = rec.get("file_name") or os.path.basename(path or "") or f"file-{pdf_id}.pdf"
    view_url = _pdf_view_url_for_path(path)
    if view_url:
        return redirect(view_url)
    if not path or not os.path.isfile(path):
        abort(404)
    return send_file(path, mimetype="application/pdf", as_attachment=False, download_name=name)

@app.route("/api/pdf_search")
def api_pdf_search():
    q = (request.args.get("q") or request.args.get("query") or "").strip()
    if not q:
        return jsonify({"ok": False, "error": "Missing query"}), 400
    rows = _pdf_db_search_by_filename(q, limit=50)
    return jsonify({"ok": True, "count": len(rows), "rows": rows})

@app.route("/api/item_overview")
def api_item_overview():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return jsonify({"ok": False, "error": _LAST_LOAD_ERR}), 503

    item = (request.args.get("item") or "").strip()
    if not item:
        abort(400, "Missing item")

    columns_so, rows_so, so_totals = _so_table_for_item(item)
    try:
        columns_po, rows_po = _po_table_for_item(item)
        open_po_cols, open_po_rows = _open_po_table_for_item(item)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    on_po_val = lookup_on_po_by_item(item)

    return jsonify(
        {
            "ok": True,
            "item": item,
            "so": {
                "columns": columns_so,
                "rows": rows_so,
                "total_on_sales": so_totals.get("on_sales_order"),
                "total_on_po": so_totals.get("on_po"),
            },
            "po": {
                "columns": columns_po,
                "rows": rows_po,
            },
            "open_po": {
                "columns": open_po_cols,
                "rows": open_po_rows,
            },
            "on_po_label": on_po_val,
        }
    )


@app.route("/so_lines")
def so_lines():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    item = (request.args.get("item") or "").strip()
    if not item:
        abort(400, "Missing item")

    columns, rows, _ = _so_table_for_item(item)

    on_po_val = lookup_on_po_by_item(item)

    return render_template_string(
        SUBPAGE_TPL,
        title=f"On Sales Order â€” {item}",
        columns=columns,
        rows=rows,
        extra_note="Source: public.wo_structured",
        on_po=on_po_val,
        open_po_columns=[],
        open_po_rows=[],
        extra_note_open_po='Source: Quickbooks',
    )

@app.route("/po_lines")
def po_lines():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    item = (request.args.get("item") or "").strip()
    if not item:
        abort(400, "Missing item")

    try:
        cols, rows = _po_table_for_item(item)
        open_cols, open_rows = _open_po_table_for_item(item)
    except ValueError as exc:
        return render_template_string(ERR_TPL, error=str(exc)), 500

    on_po_val = lookup_on_po_by_item(item)

    return render_template_string(
        SUBPAGE_TPL,
        title=f"On PO â€” {item}",
        columns=cols,
        rows=rows,
        extra_note='Source: SAP"',
        on_po=on_po_val,
        open_po_columns=open_cols,
        open_po_rows=open_rows,
        extra_note_open_po='Source: Quickbooks',
    )

@app.route("/item_details")
def item_details():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    item = (request.args.get("item") or "").strip()
    if not item:
        abort(400, "Missing item")

    columns_so, rows_so, so_totals = _so_table_for_item(item)
    try:
        columns_po, rows_po = _po_table_for_item(item)
        open_po_cols, open_po_rows = _open_po_table_for_item(item)
    except ValueError as exc:
        return render_template_string(ERR_TPL, error=str(exc)), 500

    on_po_val = lookup_on_po_by_item(item)

    return render_template_string(
        ITEM_TPL,
        item=item,
        on_po=on_po_val,
        so_columns=columns_so,
        so_rows=rows_so,
        po_columns=columns_po,
        po_rows=rows_po,
        open_po_columns=open_po_cols,
        open_po_rows=open_po_rows,
        extra_note_so="Source: public.wo_structured",
        extra_note_po='Source: SAP"',
        extra_note_open_po='Source: Quickbooks',
        so_total_on_sales=so_totals.get("on_sales_order"),
        so_total_on_po=so_totals.get("on_po"),
    )


@app.route("/item_info")
def item_info():
    try:
        df = _load_item_info(force=request.args.get("reload") == "1")
    except Exception as exc:
        return render_template_string(ERR_TPL, error=f"Item Info query failed: {exc}"), 500

    q = (request.values.get("q") or "").strip()
    rows: list[dict] = []
    display_columns = ["Part Name", "Preferred Vendor", "Original MPN", "Location", "Photo File", "Notes", "HS code", "HTS code", "Country of Origin", "Description"]
    columns = [col for col in display_columns if df is not None and col in df.columns]
    count = 0

    if q and df is not None and not df.empty:
        ql = q.lower()
        search_cols = [
            col
            for col in ("Name", "Part Name", "Description", "Type", "Taiwan MPN", "Original MPN", "Preferred Vendor")
            if col in df.columns
        ]
        if search_cols:
            mask = pd.Series(False, index=df.index)
            for col in search_cols:
                mask = mask | df[col].fillna("").astype(str).str.lower().str.contains(ql, regex=False)
            result = df.loc[mask].copy()
        else:
            result = df.iloc[0:0].copy()
        count = len(result)
        rows = result.head(200)[columns].fillna("").astype(str).to_dict(orient="records")
        for row in rows:
            if row.get("Photo File"):
                row["_photo_href"] = url_for("item_photo", path=row["Photo File"])

    return render_template_string(
        ITEM_INFO_TPL,
        loaded_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        q_val=q,
        columns=columns,
        rows=rows,
        count=count,
        shown=len(rows),
    )


@app.route("/item_photo")
def item_photo():
    photo_file = (request.args.get("path") or "").strip()
    path = _resolve_item_photo_path(photo_file)
    if path is None:
        abort(404)
    try:
        return send_file(path, as_attachment=False)
    except OSError:
        try:
            _hydrate_onedrive_file(path)
            return send_file(path, as_attachment=False)
        except OSError as exc:
            return render_template_string(
                ERR_TPL,
                error=(
                    "Photo file was found, but Windows/OneDrive would not make it available locally. "
                    f"Path: {path}. Error: {exc}"
                ),
            ), 503


def _inventory_lookup_payload(item_input: str) -> dict:
    """Keep inventory lookup/aggregation in MRP; Receiving only renders this data."""

    so_columns: list[str] | None = None
    so_rows: list[dict] | None = None
    on_hand: int | float | None = None
    on_hand_wip: int | float | None = None
    receiving_columns: list[str] = []
    receiving_rows: list[dict] = []
    inv_status_columns: list[str] = []
    inv_status_rows: list[dict] = []

    inv_filtered = INVENTORY_STATUS.copy() if INVENTORY_STATUS is not None else pd.DataFrame()
    if item_input and not inv_filtered.empty and "Part_Number" in inv_filtered.columns:
        part = inv_filtered["Part_Number"].astype(str).str.strip()
        item_norm = normalize_item(item_input)
        item_norm_upper = str(item_norm).strip().upper()
        item_upper = item_input.strip().upper()
        inv_filtered = inv_filtered.loc[
            part.str.upper().eq(item_upper) | part.str.upper().eq(item_norm_upper)
        ].copy()
    elif item_input:
        inv_filtered = pd.DataFrame()

    if not inv_filtered.empty:
        on_hand = _aggregate_metric(inv_filtered.get("On Hand", pd.Series(dtype=float)))
        on_hand_wip = _aggregate_metric(inv_filtered.get("On Hand - WIP", pd.Series(dtype=float)))

    filtered_df = SO_INV.copy()
    if item_input:
        item_norm = normalize_item(item_input)
        item_upper = item_input.strip().upper()
        item_norm_upper = str(item_norm).strip().upper()
        so_items = filtered_df["Item"].astype(str).str.strip().str.upper()
        filtered_df = filtered_df.loc[so_items.eq(item_upper) | so_items.eq(item_norm_upper)]

    if on_hand is None and not filtered_df.empty:
        on_hand, on_hand_wip = _compute_on_hand_metrics(filtered_df)

    # Resolve case/aliases before using the existing MRP lookup helpers.
    canonical_item = item_input
    if item_input:
        if not filtered_df.empty:
            canonical_item = str(filtered_df.iloc[0]['Item']).strip()
        elif not inv_filtered.empty:
            canonical_item = str(inv_filtered.iloc[0]['Part_Number']).strip()
        receiving_columns, receiving_rows = _recent_receiving_summary_for_item(item_input)
        so_columns, so_rows, _ = _so_table_for_item(canonical_item)
    else:
        so_columns, so_rows = [], []
        inv_status = INVENTORY_STATUS.copy() if INVENTORY_STATUS is not None else pd.DataFrame()
        if not inv_status.empty:
            if "Part_Number" not in inv_status.columns and "Item" in inv_status.columns:
                inv_status["Part_Number"] = inv_status["Item"]
            if "On Hand - WIP" not in inv_status.columns:
                if "In Stock(Inventory)" in inv_status.columns:
                    inv_status["On Hand - WIP"] = inv_status["In Stock(Inventory)"]
            inv_status_columns = [c for c in ("Part_Number", "On Hand", "On Hand - WIP") if c in inv_status.columns]
            if inv_status_columns:
                inv_status = inv_status[inv_status_columns].copy()
                if "Part_Number" in inv_status.columns:
                    inv_status = inv_status.sort_values("Part_Number", kind="mergesort")
                for col in ("On Hand", "On Hand - WIP"):
                    if col in inv_status.columns:
                        inv_status[col] = inv_status[col].apply(_format_intish)
                inv_status_rows = inv_status.fillna("").astype(str).to_dict(orient="records")

    return dict(
        ok=True, schema_version=1,
        loaded_at=_LAST_LOADED_AT.isoformat() if _LAST_LOADED_AT else None,
        item=item_input, canonical_item=canonical_item,
        on_hand=on_hand, on_hand_wip=on_hand_wip,
        so_columns=so_columns or [], so_rows=so_rows or [],
        receiving_columns=receiving_columns, receiving_rows=receiving_rows,
        inv_status_columns=inv_status_columns, inv_status_rows=inv_status_rows,
    )


@app.get('/api/inventory-lookup')
def api_inventory_lookup():
    global RECEIVING_LOG
    item = (request.args.get('item') or '').strip()
    if len(item) > 200:
        return jsonify(ok=False, error='Item must be at most 200 characters.'), 400
    try:
        if request.args.get('reload') == '1':
            _load_from_db(force=True)
            RECEIVING_LOG = None
        else:
            _ensure_loaded()
        if _LAST_LOAD_ERR:
            return jsonify(ok=False, error='MRP inventory data is unavailable. Check the MRP load.'), 503
        return jsonify(_inventory_lookup_payload(item))
    except Exception:
        app.logger.exception('Inventory lookup failed')
        return jsonify(ok=False, error='MRP inventory lookup failed.'), 503


@app.route("/api/global_suggest")
def api_global_suggest():
    _ensure_loaded()
    q = (request.args.get("q") or request.args.get("query") or "").strip()
    if not q:
        return jsonify({"ok": True, "items": []})
    try:
        ql = q.lower()
        rows = GLOBAL_SEARCH_INDEX
        starts = [r for r in rows if str(r.get("label", "")).lower().startswith(ql)]
        contains = [r for r in rows if ql in str(r.get("search", r.get("label", ""))).lower() and r not in starts]
        out = []
        for row in (starts + contains)[:20]:
            out.append(
                {
                    "type": row.get("type", ""),
                    "label": row.get("label", ""),
                    "href": row.get("href", ""),
                    "on_hand": row.get("on_hand", ""),
                    "customer": row.get("customer", ""),
                    "ship_date": row.get("ship_date", ""),
                }
            )
        return jsonify({"ok": True, "items": out})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

@app.route("/global_search")
def global_search():
    _ensure_loaded()
    q = (request.args.get("q") or "").strip()
    if not q:
        return redirect(url_for("index"))
    ql = q.lower()
    rows = GLOBAL_SEARCH_INDEX
    exact = next((r for r in rows if str(r.get("label", "")).lower() == ql), None)
    if exact and exact.get("href"):
        return redirect(str(exact["href"]))
    starts = next((r for r in rows if str(r.get("label", "")).lower().startswith(ql)), None)
    if starts and starts.get("href"):
        return redirect(str(starts["href"]))
    if q.upper().startswith("SO") or q.replace("-", "").isdigit():
        return redirect(url_for("index", so=q))
    return redirect(url_for("index", customer=q))

@app.route("/api/item_suggest")
def api_item_suggest():
    _ensure_loaded()
    q = (request.args.get("q") or request.args.get("query") or "").strip()
    if not q:
        return jsonify({"ok": True, "items": []})
    try:
        items = ITEM_SUGGEST_CACHE
        ql = q.lower()
        starts = [i for i in items if i.lower().startswith(ql)]
        contains = [i for i in items if ql in i.lower() and i not in starts]
        out = (starts + contains)[:20]
        return jsonify({"ok": True, "items": out})
    except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 500


@app.route("/api/item_info_suggest")
def api_item_info_suggest():
    q = (request.args.get("q") or request.args.get("query") or "").strip()
    if not q:
        return jsonify({"ok": True, "items": []})
    try:
        _load_item_info()
        items = ITEM_INFO_SUGGEST_CACHE
        ql = q.lower()
        starts = [i for i in items if i.lower().startswith(ql)]
        contains = [i for i in items if ql in i.lower() and i not in starts]
        out = [
            {"text": item, "has_photo": ITEM_INFO_PHOTO_BY_SUGGESTION.get(item, False)}
            for item in (starts + contains)[:20]
        ]
        return jsonify({"ok": True, "items": out})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


@app.route("/api/quotation_item_suggest")
def api_quotation_item_suggest():
    _ensure_loaded()
    q = (request.args.get("q") or request.args.get("query") or "").strip()
    if not q:
        return jsonify({"ok": True, "items": []})
    try:
        rows = QUOTE_ITEM_SUGGEST_ROWS
        ql = q.lower()
        starts = [r for r in rows if str(r.get("item", "")).lower().startswith(ql)]
        contains = [r for r in rows if ql in str(r.get("item", "")).lower() and r not in starts]
        out = (starts + contains)[:20]
        return jsonify({"ok": True, "items": out})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

def _peripheral_workbook_path() -> Path | None:
    configured = os.environ.get("PERIPHERAL_STATUS_WORKBOOK", "").strip()
    if configured:
        candidate = Path(configured).expanduser()
        return candidate if candidate.is_file() else None
    configured_path = Path(PERIPHERAL_STATUS_FILE)
    return configured_path if configured_path.is_file() else None


def _display_excel_value(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _excel_fill_class(cell) -> str:
    if not cell.fill or cell.fill.fill_type is None:
        return ""
    color = cell.fill.fgColor
    rgb = str(color.rgb or "").upper()[-6:] if color.type == "rgb" else ""
    indexed = color.indexed if color.type == "indexed" else None
    if rgb in {"FFFF00", "FFF2CC", "FFE699", "FFD966", "FFFF99"} or indexed in {6, 13}:
        return "model-recommended"
    if rgb in {"FF0000", "FFC7CE", "F4CCCC", "EA9999", "E06666"} or indexed in {2, 10}:
        return "model-do-not-use"
    return ""


def _load_peripheral_status(force: bool = False) -> dict:
    global PERIPHERAL_STATUS_CACHE, PERIPHERAL_STATUS_CACHE_KEY
    path = _peripheral_workbook_path()
    if path is None:
        raise FileNotFoundError("Place 'Peripheral Status Update_YYYYMMDD.xlsx' in the MRP_System base folder, or set PERIPHERAL_STATUS_WORKBOOK to its full path.")
    stat = path.stat()
    cache_key = (str(path.resolve()).lower(), stat.st_mtime_ns, stat.st_size)
    if not force and PERIPHERAL_STATUS_CACHE is not None and PERIPHERAL_STATUS_CACHE_KEY == cache_key:
        return PERIPHERAL_STATUS_CACHE
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        names, sheets, missing = list(workbook.sheetnames), [], []
        for label, aliases in (("SSD", ("ssd",)), ("DDR / Memory", ("ddr", "memory"))):
            sheet_name = next((n for n in names if any(a in n.lower() for a in aliases)), None)
            if sheet_name is None:
                missing.append(label)
                continue
            ws = workbook[sheet_name]
            values = [list(row) for row in ws.iter_rows(min_col=1, max_col=16)]
            while values and not any(cell.value is not None for cell in values[-1]):
                values.pop()
            if not values:
                sheets.append({"label": label, "headers": [chr(65+i) for i in range(16)], "rows": []})
                continue
            header_idx = next((i for i, row in enumerate(values) if any(c.value is not None for c in row)), 0)
            source_headers = [_display_excel_value(c).strip() or chr(65+i) for i, c in enumerate(values[header_idx])]
            hidden_headers = {"M.2 SSD", "常備料", "HQ Status", "WH01S"}
            if label == "DDR / Memory":
                hidden_headers.update({"DDR4-3200", "NTA Available", "Sales/week", "Category RR"})
            visible_indexes = [
                i for i, header in enumerate(source_headers)
                if i != 1 and header not in hidden_headers
            ]
            headers = [source_headers[i] for i in visible_indexes]
            rows = []
            for excel_row in values[header_idx + 1:]:
                source_cells = [_display_excel_value(c) for c in excel_row]
                cells = [source_cells[i] for i in visible_indexes]
                if any(value.strip() for value in source_cells):
                    rows.append({
                        "cells": cells,
                        "model": source_cells[4].strip(),
                        "status": source_cells[6].strip(),
                        "cost": source_cells[7].strip(),
                        "selling_price": source_cells[8].strip(),
                        "hq_available": source_cells[9].strip(),
                        "model_class": _excel_fill_class(excel_row[4]),
                        "model_display_index": visible_indexes.index(4) if 4 in visible_indexes else -1,
                        "search_text": " ".join(source_cells).lower(),
                    })
            sheets.append({"label": label, "headers": headers, "rows": rows})
    finally:
        workbook.close()
    result = {"sheets": sheets, "warnings": (["Missing sheet(s): " + ", ".join(missing) + "."] if missing else []),
              "workbook_name": path.name, "loaded_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    PERIPHERAL_STATUS_CACHE, PERIPHERAL_STATUS_CACHE_KEY = result, cache_key
    return result


@app.route("/quotation_lookup/peripheral_status")
def peripheral_status():
    try:
        context = _load_peripheral_status(force=request.args.get("reload") == "1")
        return render_template_string(PERIPHERAL_STATUS_TPL, error=None, **context)
    except Exception as exc:
        return render_template_string(PERIPHERAL_STATUS_TPL, error=str(exc), sheets=[], warnings=[],
                                      workbook_name="", loaded_at="")


@app.route("/quotation_lookup")
def quotation_lookup():
    _ensure_loaded()
    if _LAST_LOAD_ERR:
        return render_template_string(ERR_TPL, error=_LAST_LOAD_ERR), 503

    item_input = (request.values.get("item") or "").strip()
    item_lookup = _resolve_ledger_item_key(item_input)
    item_cards, companion_error = load_item_cards(item_lookup or item_input, QUOTE_ITEM_SUGGEST_ROWS)
    qty_val = 1

    ledger_columns: list[str] = []
    ledger_rows: list[dict] = []
    opening_qty = None
    earliest_atp = None
    cache_key = (item_lookup, 1)
    cached = QUOTATION_VIEW_CACHE.get(cache_key)
    if cached is not None:
        ledger_columns = cached.get("ledger_columns", [])
        ledger_rows = cached.get("ledger_rows", [])
        opening_qty = cached.get("opening_qty")
        earliest_atp = cached.get("earliest_atp")
        return render_template_string(
            QUOTE_TPL,
            item_cards=item_cards,
            companion_error=companion_error,
            item_val=item_input,
            opening_qty=opening_qty,
            earliest_atp=earliest_atp,
            ledger_columns=ledger_columns,
            ledger_rows=ledger_rows,
            loaded_at=_LAST_LOADED_AT.strftime("%Y-%m-%d %H:%M:%S") if _LAST_LOADED_AT else "",
        )

    if item_lookup and LEDGER is not None and not LEDGER.empty:
        if item_lookup in LEDGER_ITEM_INDEX:
            df_item = LEDGER_ITEM_INDEX[item_lookup].copy()
        else:
            df = LEDGER
            item_col = "Item" if "Item" in df.columns else "Item_raw"
            df_item = df.loc[df[item_col].astype(str).str.strip().str.upper() == item_lookup].copy()
        if not df_item.empty:
            # Opening snapshot:
            # 1) Prefer explicit OPEN rows; 2) if none, fall back to any Opening values.
            if "Opening" in df_item.columns:
                open_rows = df_item.loc[df_item["Kind"].astype(str) == "OPEN"].copy()
                if not open_rows.empty:
                    open_rows = open_rows.sort_values("Date")
                    opening_qty = _aggregate_metric(open_rows["Opening"])
                else:
                    opening_series = pd.to_numeric(df_item["Opening"], errors="coerce").dropna()
                    if not opening_series.empty:
                        opening_qty = _aggregate_metric(opening_series)

            # Prepare ledger table rows (date, kind, delta, projected)
            keep_cols: list[str] = []
            for c in (
                "Date",
                "Kind",
                "Source",
                "Delta",
                "Projected_NAV",
                "NAV_before",
                "NAV_after",
                "QB Num",
                "P. O. #",
                "Name",
            ):
                if c in df_item.columns and c not in keep_cols:
                    keep_cols.append(c)

            if keep_cols:
                # Preferred column order: Date, Name, QB Num, then others, Waiting_Item last
                preferred: list[str] = []
                if "Date" in keep_cols:
                    preferred.append("Date")
                if "Name" in keep_cols:
                    preferred.append("Name")
                if "QB Num" in keep_cols:
                    preferred.append("QB Num")
                rest = [c for c in keep_cols if c not in preferred and c != "Waiting_Item"]
                keep_cols = preferred + rest

                df_item = df_item.sort_values(["Date", "Kind"])
                date_vals = pd.to_datetime(df_item["Date"], errors="coerce")
                df_item["Date"] = date_vals.dt.strftime("%Y-%m-%d")

                # Determine minimum Projected_NAV for highlighting (after sort so alignment is correct)
                min_nav_value = None
                proj_series = None
                if "Projected_NAV" in df_item.columns:
                    proj_series = pd.to_numeric(df_item["Projected_NAV"], errors="coerce")
                    if not proj_series.dropna().empty:
                        min_nav_value = proj_series.min()

                if "Projected_NAV" in df_item.columns:
                    df_item.rename(columns={"Projected_NAV": "Projected_Qty"}, inplace=True)
                    keep_cols = ["Projected_Qty" if c == "Projected_NAV" else c for c in keep_cols]

                for c in ("Delta", "Projected_NAV", "NAV_before", "NAV_after"):
                    if c in df_item.columns:
                        df_item[c] = df_item[c].apply(_format_intish)
                if "Projected_Qty" in df_item.columns:
                    df_item["Projected_Qty"] = df_item["Projected_Qty"].apply(_format_intish)

                records = df_item[keep_cols].fillna("").astype(str).to_dict(orient="records")

                if "QB Num" in keep_cols and "Waiting_Item" not in keep_cols:
                    keep_cols.append("Waiting_Item")
                    for rec in records:
                        qb = rec.get("QB Num", "")
                        rec["Waiting_Item"] = WAITING_ITEMS_BY_QB.get(str(qb), "{}")

                # Attach _is_min_nav flag for UI highlighting
                if min_nav_value is not None and proj_series is not None:
                    proj_vals = proj_series.tolist()
                    for rec, v in zip(records, proj_vals):
                        try:
                            fv = float(v)
                        except Exception:
                            fv = None
                        rec["_is_min_nav"] = (fv == float(min_nav_value))
                else:
                    for rec in records:
                        rec["_is_min_nav"] = False

                ledger_columns = keep_cols
                ledger_rows = records


        earliest_atp_dt = _lookup_earliest_atp_date(item_lookup, qty=qty_val)
        if earliest_atp_dt is not None:
            earliest_atp = earliest_atp_dt.strftime("%Y-%m-%d")
        else:
            earliest_atp = "Out of Stock"

    if item_lookup:
        if len(QUOTATION_VIEW_CACHE) >= 256:
            QUOTATION_VIEW_CACHE.clear()
        QUOTATION_VIEW_CACHE[cache_key] = {
            "ledger_columns": ledger_columns,
            "ledger_rows": ledger_rows,
            "opening_qty": opening_qty,
            "earliest_atp": earliest_atp,
        }

    return render_template_string(
        QUOTE_TPL,
        item_cards=item_cards,
        companion_error=companion_error,
        item_val=item_lookup or item_input,
        opening_qty=opening_qty,
        earliest_atp=earliest_atp,
        ledger_columns=ledger_columns,
        ledger_rows=ledger_rows,
        loaded_at=_LAST_LOADED_AT.strftime("%Y-%m-%d %H:%M:%S") if _LAST_LOADED_AT else "",
    )


if __name__ == "__main__":
    # Flask dev server
    # Preload PDF map on startup for faster first-hit
    _load_pdf_map(force=True)
    app.run(debug=True, host="0.0.0.0", port=5002)

