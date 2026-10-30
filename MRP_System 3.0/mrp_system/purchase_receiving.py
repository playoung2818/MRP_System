"""POD-scoped, previewed, append-only receipts into the EXISTING receiving_log.

No ORM startup/schema creation. Shares the receiving app's table lock and workbook
format, but deliberately offers no history replacement or deletion actions.
"""
from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import math
import re
import secrets
from zipfile import ZipFile, BadZipFile

from flask import abort, current_app, render_template_string, request, session
from itsdangerous import BadSignature, URLSafeTimedSerializer
from openpyxl import load_workbook
from sqlalchemy import MetaData, Table, select, text, func
from sqlalchemy.exc import SQLAlchemyError
from werkzeug.exceptions import RequestEntityTooLarge, HTTPException

from mrp_system.normalize.mrp_normalize import normalize_item
from mrp_system.runtime.db_config import get_engine

HEADERS = {'Date':'entry_date', 'Inv#':'invoice_number', 'Box #':'box_number',
           'POD#':'pod_number', 'Part#':'part_number', 'SN#':'serial_number', 'QTY':'quantity'}
KEYS = tuple(HEADERS.values())
MAX_UPLOAD = 10 * 1024 * 1024
WORKBOOK_VENDOR = 'Neousys Technology Incorp.'


def clean(value):
    return '' if value is None else str(value).strip()


def pod_id(value):
    value = re.sub(r'^POD[ #:-]*', '', clean(value).upper())
    if not re.fullmatch(r'\d{6,10}', value):
        raise ValueError('Enter a POD number such as POD-261003 or 261003.')
    return 'POD-' + value


def vendor_label(source):
    return ' / '.join(sorted({clean(row.get('Source Name')) for row in source
                              if clean(row.get('Source Name'))},key=str.casefold))


def validate_cpu_batches(rows):
    for row in rows:
        if clean(row['part_number']).casefold().startswith(('i','core')):
            serial = clean(row['serial_number'])
            if serial.upper() != 'NA':
                pieces = serial.rsplit('_',1)
                if len(pieces) != 2 or not all(piece.strip() for piece in pieces):
                    raise ValueError('Items starting with I or Core require both serial number and batch number (serial_batch).')


def fingerprint(rows):
    return sha256(json.dumps(sorted(json.dumps(dict(r), sort_keys=True, default=str) for r in rows)).encode()).hexdigest()


def source_context(conn, pod):
    rows = [dict(r) for r in conn.execute(text('''SELECT * FROM public."Open_Purchase_Orders"
        WHERE UPPER(BTRIM(COALESCE(NULLIF(BTRIM("POD#"), ''), "QB Num"))) IN (:pod,:digits,:compact)
        ORDER BY "Item", "Inventory Site", ctid'''),
        {'pod':pod,'digits':pod[4:],'compact':pod.replace('-','')}).mappings()]
    if not rows:
        raise ValueError('This POD is not in the current open purchase orders. Receiving is unavailable.')
    master = {str(name).casefold():name for name in conn.execute(text(
        'SELECT "Part_Number" FROM public.inventory_status')).scalars() if name}
    parts = {}
    for row in rows:
        raw = clean(row.get('Item'))
        if not raw:
            continue
        normalized = normalize_item(raw)
        canonical = master.get(normalized.casefold(), master.get(raw.casefold(), raw))
        parts[raw.casefold()] = canonical
        parts[normalized.casefold()] = canonical
        parts[canonical.casefold()] = canonical
        row['display_item'] = canonical
    # Presentation-only labels must not become part of the source concurrency token.
    source_hash = fingerprint([{k:v for k,v in row.items() if k != 'display_item'} for row in rows])
    return rows, parts, source_hash


def upload_rows(content):
    if len(content) > MAX_UPLOAD:
        raise ValueError('Upload must be smaller than 10 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Workbook is too large after decompression.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('Upload a valid .xlsx or .xlsm workbook.') from exc
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean(v).casefold() for v in next(iterator, ())]
        if any(headers.count(key.casefold()) != 1 for key in HEADERS):
            raise ValueError('First worksheet headers must include exactly one each: ' + ', '.join(HEADERS))
        positions = {key:headers.index(header.casefold()) for header,key in HEADERS.items()}
        rows = []
        for index, values in enumerate(iterator, 2):
            if index > 20001:
                raise ValueError('Upload at most 20,000 worksheet rows.')
            rows.append({key:values[pos] if pos < len(values) else None for key,pos in positions.items()})
        return rows
    finally:
        workbook.close()


def manual_rows(form, pod, warnings=None):
    parts, serials, quantities = (form.getlist(key) for key in ('part_number','serial_number','quantity'))
    if not len(parts) == len(serials) == len(quantities) or len(parts) > 1000:
        raise ValueError('Invalid receiving rows; enter at most 1,000 rows.')
    result = []
    for index, (part, serial_text, quantity) in enumerate(zip(parts,serials,quantities),1):
        try:
            total = float(quantity)
            if not math.isfinite(total) or total <= 0:
                raise ValueError()
        except (ValueError,TypeError,OverflowError):
            raise ValueError(f'Row {index}: QTY must be finite and positive.') from None
        serial_lines = [clean(serial) for serial in serial_text.splitlines() if clean(serial)]
        all_na = bool(serial_lines) and all(serial.upper() == 'NA' for serial in serial_lines)
        if all_na:
            # Repeated NA placeholders still describe one non-serialized batch.
            serial_lines = ['NA']
        elif serial_lines and len(serial_lines) != total and warnings is not None:
            warnings.append(f'Row {index} ({part}): {len(serial_lines)} serial-number lines do not match QTY {total:g}. '
                            f'Preview records one unit per serial ({len(serial_lines)} units), not QTY multiplied by serial count. '
                            'Check and correct the entries before confirming.')
        for serial in serial_lines or ['']:
            result.append(dict(entry_date=form.get('entry_date'),invoice_number=form.get('invoice_number'),
                               box_number=form.get('box_number'),pod_number=pod,part_number=part,
                               serial_number=serial,quantity=total if all_na else 1.0))
    return result


def validate_rows(rows, pod, parts):
    if len(rows) > 20000:
        raise ValueError('Enter at most 20,000 receipt events.')
    result, skipped, errors = [], 0, []
    for index, raw in enumerate(rows, 1):
        row = {key:clean(raw.get(key)) for key in KEYS if key not in ('entry_date','quantity')}
        if row['serial_number'].upper() in ('','NAN'):
            skipped += 1
            continue
        try:
            if any(len(value) > 255 or value.startswith('=') for value in row.values()):
                raise ValueError('Text fields must be at most 255 characters and not formulas.')
            if pod_id(row['pod_number']) != pod:
                raise ValueError('Workbook contains another POD. Upload only receipts for this POD.')
            row['pod_number'] = pod
            raw_part = row['part_number']
            canonical = parts.get(raw_part.casefold()) or parts.get(normalize_item(raw_part).casefold())
            if not canonical:
                raise ValueError('Part is not on this POD: ' + raw_part)
            row['part_number'] = canonical
            day = raw.get('entry_date')
            if isinstance(day, datetime):day = day.date()
            if not isinstance(day, date):day = date.fromisoformat(clean(day))
            row['entry_date'] = day.isoformat()
            quantity = 1.0 if clean(raw.get('quantity')) == '' else float(raw['quantity'])
            if not math.isfinite(quantity) or quantity <= 0:
                raise ValueError('Quantity must be finite and positive.')
            row['quantity'] = quantity
            result.append(row)
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append(f'Row {index}: {exc}')
    if errors:
        raise ValueError('Nothing saved. ' + ' '.join(errors[:20]))
    if not result:
        raise ValueError('Enter serial numbers, or NA for non-serialized receipts.')
    return result, skipped


def signature(row):
    day = row['entry_date']
    if isinstance(day, datetime):day = day.date()
    pod = clean(row['pod_number'])
    try:pod = pod_id(pod)
    except ValueError:pass
    return (str(day), clean(row['invoice_number']), clean(row['box_number']), pod,
            clean(row['part_number']), clean(row['serial_number']).upper(), float(row['quantity']))


def receipt_key(row):
    serial = clean(row['serial_number']).upper()
    return 'NA:' + sha256(json.dumps(signature(row)).encode()).hexdigest() if serial == 'NA' else serial


def receiving_table(conn):
    table = Table('receiving_log', MetaData(), schema='public', autoload_with=conn)
    if not set((*KEYS,'id')).issubset(table.c.keys()):
        raise ValueError('The existing receiving_log table is missing required columns.')
    return table


def database_groups(conn, table, keys):
    groups = defaultdict(list)
    serials = sorted({'NA' if key.startswith('NA:') else key for key in keys})
    query = select(table).where(func.upper(func.btrim(table.c.serial_number)).in_(serials))
    for row in conn.execute(query).mappings():
        groups[receipt_key(row)].append(dict(row))
    return groups


def preview(conn, rows):
    existing = database_groups(conn, receiving_table(conn), {receipt_key(row) for row in rows})
    uploaded, skipped = defaultdict(list), 0
    for row in rows:
        key = receipt_key(row)
        if any(signature(row) == signature(other) for other in uploaded[key]):
            skipped += 1
        else:uploaded[key].append(row)
    review = []
    for key, candidates in uploaded.items():
        matches = existing.get(key, [])
        known = {signature(row) for row in matches}
        new = [row for row in candidates if signature(row) not in known]
        skipped += len(candidates)-len(new)
        if new:
            review.append(dict(key=key,uploaded=new,existing=matches,snapshot=fingerprint(matches),
                               kind='new' if not matches and len(new)==1 else 'conflict'))
    return review, skipped


def apply_receipts(engine, pod, source_hash, review, choices):
    inserted, skipped = 0, 0
    with engine.connect() as conn:
        try:
            if conn.dialect.name != 'postgresql':
                raise ValueError('Safe receiving requires the configured PostgreSQL database.')
            conn.execute(text("SET LOCAL lock_timeout = '10s'"))
            # Same lock as port 5000: concurrent receivers cannot race serial checks.
            conn.execute(text('LOCK TABLE public.receiving_log IN SHARE ROW EXCLUSIVE MODE'))
            conn.execute(text('LOCK TABLE public."Open_Purchase_Orders" IN SHARE MODE'))
            source, parts, current_source = source_context(conn, pod)
            if current_source != source_hash:
                raise ValueError('POD changed since preview. Nothing saved; preview again.')
            table = receiving_table(conn)
            current = database_groups(conn, table, {group['key'] for group in review})
            pending = []
            for index, group in enumerate(review):
                choice = choices.get(str(index), 'skip')
                if choice not in ('add','skip'):
                    raise ValueError('Invalid receiving choice.')
                if choice == 'skip':
                    skipped += len(group['uploaded'])
                    continue
                rows, _ = validate_rows(group['uploaded'], pod, parts)
                if any(clean(line.get('Source Name')) != WORKBOOK_VENDOR for line in source):
                    validate_cpu_batches(rows)
                    if any(row['box_number'] != '1--1' for row in rows):
                        raise ValueError('Box # must be 1--1 for this vendor. Preview again.')
                    if any(row['invoice_number'] != vendor_label(source) for row in rows):
                        raise ValueError('Invoice # must be the vendor name. Preview again.')
                matches = current.get(group['key'], [])
                known = {signature(row) for row in matches}
                new = [row for row in rows if signature(row) not in known]
                if fingerprint(matches) != group['snapshot'] and new:
                    raise ValueError('Receiving records changed since preview. Nothing saved; preview again.')
                skipped += len(rows)-len(new)
                for row in new:
                    pending.append(dict(row,entry_date=date.fromisoformat(row['entry_date'])))
                    known.add(signature(row))
            if pending:
                conn.execute(table.insert(), pending)
                inserted = len(pending)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    return inserted, skipped


def register_receiving_routes(bp):
    @bp.route('/purchase_orders/pod/<path:pod_number>', methods=['GET','POST'])
    def pod_detail(pod_number):
        from mrp_system.purchase_order_lookup import display_value
        request.max_content_length = MAX_UPLOAD
        try:pod = pod_id(pod_number)
        except ValueError as exc:return str(exc), 400
        csrf = session.setdefault('pod_receiving_csrf', secrets.token_urlsafe(32))
        context = dict(pod=pod,csrf=csrf,today=date.today().isoformat(),review=None,
                       error=None,message=None,rows=[],parts=[],receipts=[],fields=KEYS,
                       canonical_item=current_app.jinja_env.filters.get('canonical_item',clean),
                       display_value=display_value,manual=[],allow_upload=False,vendor_name='',warnings=[])
        engine, status = get_engine(), 200
        signer = URLSafeTimedSerializer(current_app.secret_key,salt='pod-receiving-v1')
        try:
            if request.method == 'POST' and not secrets.compare_digest(request.form.get('csrf_token',''),csrf):
                abort(400)
            if request.method == 'POST' and request.form.get('action') == 'manual':
                context['manual'] = [dict(part_number=p,serial_number=s,quantity=q) for p,s,q in zip(
                    request.form.getlist('part_number'),request.form.getlist('serial_number'),request.form.getlist('quantity'))]
            with engine.connect() as conn:
                source, parts, source_hash = source_context(conn,pod)
                context.update(rows=source,parts=sorted(set(parts.values()),key=str.casefold),
                               vendor_name=vendor_label(source),allow_upload={clean(row.get('Source Name')) for row in source} == {WORKBOOK_VENDOR})
                if 'canonical_item' not in current_app.jinja_env.filters:
                    context['canonical_item'] = lambda value: parts.get(clean(value).casefold(),value)
                # Source and received quantities are separate facts, not additive balances.
                context['receipts'] = list(conn.execute(text('''SELECT entry_date,invoice_number,box_number,
                    part_number,serial_number,quantity FROM public.receiving_log
                    WHERE UPPER(BTRIM(pod_number)) IN (:pod,:digits,:compact)
                    ORDER BY entry_date DESC,id DESC LIMIT 500'''),
                    {'pod':pod,'digits':pod[4:],'compact':pod.replace('-','')}).mappings())
                if request.method == 'POST' and request.form.get('action') != 'confirm':
                    if request.form.get('action') == 'upload':
                        if not context['allow_upload']:
                            raise ValueError('Workbook upload is available only for Neousys Technology Incorp. PODs. Use manual receiving.')
                        file = request.files.get('file')
                        if not file or not file.filename.lower().endswith(('.xlsx','.xlsm')):
                            raise ValueError('Choose an .xlsx or .xlsm workbook.')
                        raw = upload_rows(file.read(MAX_UPLOAD+1))
                    elif request.form.get('action') == 'manual':
                        raw = manual_rows(request.form,pod,context['warnings'])
                        if not context['allow_upload']:
                            raw = [dict(row,box_number='1--1',invoice_number=context['vendor_name']) for row in raw]
                    else:raise ValueError('Choose a receiving action.')
                    records, blanks = validate_rows(raw,pod,parts)
                    if not context['allow_upload']:
                        validate_cpu_batches(records)
                    review, duplicates = preview(conn,records)
                    context.update(review=review,skipped=blanks+duplicates)
                    if review:
                        signed = [{k:v for k,v in group.items() if k != 'existing'} for group in review]
                        context['token'] = signer.dumps(dict(version=1,csrf=csrf,pod=pod,source=source_hash,review=signed,
                                                            entry_method=request.form.get('action'),warnings=context['warnings']))
                    else:context['message']='Everything already matches. Nothing saved.'
            if request.method == 'POST' and request.form.get('action') == 'confirm':
                try:payload = signer.loads(request.form.get('preview_token',''),max_age=1800)
                except BadSignature as exc:raise ValueError('Preview expired or invalid. Preview again.') from exc
                if payload.get('version') != 1 or payload.get('pod') != pod or payload.get('csrf') != csrf:
                    raise ValueError('Preview does not belong to this POD/session.')
                if payload.get('entry_method') == 'upload' and not context['allow_upload']:
                    raise ValueError('Workbook upload is unavailable for this vendor. Use manual receiving.')
                inserted, skipped = apply_receipts(engine,pod,payload['source'],payload['review'],
                    {str(i):request.form.get(f'choice_{i}','skip') for i in range(len(payload['review']))})
                # Reload after commit so the newly appended receipts are immediately visible.
                from flask import redirect, url_for, flash
                flash(f'Receiving saved: {inserted} receipt rows added; {skipped} skipped. Existing history unchanged.',
                      'pod-receiving')
                return redirect(url_for('purchase_orders.pod_detail',pod_number=pod),code=303)
        except RequestEntityTooLarge:
            context['error'], status = 'Upload must be smaller than 10 MB.', 413
        except ValueError as exc:
            context['error'], status = str(exc), 400
        except SQLAlchemyError:
            current_app.logger.exception('POD receiving failed')
            context['error'], status = 'Database operation failed. Nothing saved; retry with a fresh preview.', 503
        except HTTPException:
            raise
        except Exception:
            current_app.logger.exception('POD receiving input processing failed')
            context['error'], status = 'Unable to process this input. Nothing saved; check the workbook or entries.', 400
        finally:
            engine.dispose()
        return render_template_string(TEMPLATE,**context),status


TEMPLATE = '''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{{ pod }} · Purchase receiving</title><link rel="icon" href="/static/robot.svg" type="image/svg+xml">
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<link href="/static/workspace.css?v=2" rel="stylesheet"><style>
.receipt-serial-line{display:grid;grid-template-columns:24px minmax(160px,1fr);gap:6px;margin-bottom:6px;align-items:center}.receipt-serial-line.has-batch{grid-template-columns:24px minmax(160px,1fr) minmax(140px,1fr)}.receipt-serial-line[hidden],.receipt-serial-panel[hidden],.receipt-serial-line [hidden]{display:none}.receipt-serial-panel{min-width:230px;max-height:440px;overflow:auto;padding:2px}.receipt-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}.receipt-table{min-width:750px}.receipt-table>tbody>tr>td{vertical-align:top!important}.receipt-table textarea{min-width:200px}.receipt-table select{min-width:220px}.receipt-table input{min-width:100px}.receipt-card{margin-bottom:18px}.receipt-summary{padding:12px;background:#f5f8fc}.receipt-history{max-height:380px;overflow:auto}@media(max-width:650px){.receipt-grid{grid-template-columns:1fr}}
</style></head><body class="workspace-page"><header class="workspace-nav"><div class="workspace-nav-inner"><a class="workspace-brand" href="/">NTA MRP</a><a class="workspace-home" href="/purchase_orders">Purchase Orders</a></div></header>
<main class="workspace-main"><div class="workspace-breadcrumb"><a href="/purchase_orders">Purchase Orders</a><span>/</span><span>{{ pod }}</span></div>
<header class="workspace-header"><div><h1 class="workspace-title">{{ pod }}</h1><p class="workspace-description">Receive parts, review serials, and append receipts without replacing history.</p></div></header>
<form class="workspace-search d-flex gap-2" action="/purchase_orders" method="get"><label class="visually-hidden" for="pod-search">Search POD#</label><input class="form-control" id="pod-search" name="pod_num" type="search" placeholder="Search POD#" autocomplete="off"><button class="btn btn-outline-primary">Search</button></form>
{% for message in get_flashed_messages(category_filter=['pod-receiving']) %}<div class="alert alert-success" role="status">{{ message }}</div>{% endfor %}
{% if error %}<div class="alert alert-danger" role="alert">{{ error }}</div>{% endif %}{% if message %}<div class="alert alert-success" role="status">{{ message }}</div>{% endif %}
{% if rows %}<section class="card-lite bg-white receipt-card"><div class="card-header"><strong>Current POD lines</strong></div><div class="table-responsive"><table class="table table-sm"><thead><tr><th>Part</th><th>Vendor</th><th>Source Qty</th><th>Source received</th><th>Source outstanding</th><th>Site</th><th>Delivery</th></tr></thead><tbody>{% for row in rows %}<tr><td>{{ row.display_item }}</td><td>{{ row['Source Name'] or '' }}</td><td>{{ display_value(row['Qty']) }}</td><td>{{ display_value(row["Rcv'd"]) }}</td><td>{{ display_value(row['Qty(+)']) }}</td><td>{{ row['Inventory Site'] or '' }}</td><td>{{ display_value(row['Deliv Date']) }}</td></tr>{% endfor %}</tbody></table></div></section>
<p class="small text-muted">Source quantities come from the latest PO import. Receiving here writes only receiving_log; it does not update QuickBooks PO quantities or immediately change inventory balances. Receipt quantities are not subtracted again from source outstanding.</p>
<section class="card-lite bg-white receipt-card"><div class="card-header"><strong>Receive this POD</strong></div><div class="p-3"><form method="post" id="manual-receiving" data-expanded="{{ 'false' if allow_upload else 'true' }}"><input type="hidden" name="csrf_token" value="{{ csrf }}"><input type="hidden" name="action" value="manual">
<div class="receipt-grid mb-3"><div><label for="receipt-date" class="form-label">Received date</label><input class="form-control" id="receipt-date" type="date" name="entry_date" value="{{ request.form.get('entry_date', today) if manual else today }}" required></div><div><label for="receipt-invoice" class="form-label">Invoice #</label><input class="form-control" id="receipt-invoice" name="invoice_number" maxlength="255" value="{{ vendor_name if not allow_upload else (request.form.get('invoice_number','') if manual else '') }}" {{ 'readonly' if not allow_upload else '' }}></div><div><label for="receipt-box" class="form-label">Box #</label><input class="form-control" id="receipt-box" name="box_number" maxlength="255" value="{{ '1--1' if not allow_upload else (request.form.get('box_number','') if manual else '') }}" {{ 'readonly' if not allow_upload else '' }}></div></div>
<p class="small text-muted">{% if allow_upload %}Enter total QTY and one serial per line (one unit per serial).{% else %}QTY creates one serial-entry row per unit (except items starting with Dust). Items starting with I or Core require a serial and batch number, saved as serial_batch. Press Enter to move from serial to batch, then to the next serial.{% endif %} Preview warns if serial count differs from QTY. For non-serialized stock, enter NA; all-NA entries use total QTY and skip the count check.</p>
<div class="table-responsive"><table class="table receipt-table"><thead><tr><th>Part</th><th>Serial numbers</th><th>QTY</th><th></th></tr></thead><tbody id="receipt-lines"></tbody></table></div>
<button class="btn btn-outline-secondary" type="button" id="add-receipt">Add row</button> <button class="btn btn-primary">Preview receiving</button></form></div></section>
{% if allow_upload %}<section class="card-lite bg-white receipt-card"><div class="card-header"><strong>Or upload receiving workbook</strong></div><div class="p-3"><p class="small">Same format as port 5000: first worksheet, headers <code>Date, Inv#, Box #, POD#, Part#, SN#, QTY</code>. Only this POD's parts are allowed. Limit 10 MB / 20,000 rows. Blank/nan serials are skipped; blank quantity defaults to 1.</p><form method="post" enctype="multipart/form-data"><input type="hidden" name="csrf_token" value="{{ csrf }}"><input type="hidden" name="action" value="upload"><label class="visually-hidden" for="receipt-file">Receiving workbook</label><input class="form-control mb-2" id="receipt-file" type="file" name="file" accept=".xlsx,.xlsm" required><button class="btn btn-primary">Preview workbook</button></form></div></section>{% endif %}
{% endif %}
{% for warning in warnings %}<div class="alert alert-warning" role="alert">{{ warning }}</div>{% endfor %}
{% if review %}<section class="card-lite bg-white receipt-card"><div class="card-header"><strong>Receiving preview — nothing saved yet</strong></div><div class="p-3"><p>{{ review|length }} serial/event groups; {{ skipped }} blank or exact repeat rows skipped. Existing serial conflicts default to Skip. Adding a new event preserves every existing row. No update/delete actions are available.</p><form method="post"><input type="hidden" name="csrf_token" value="{{ csrf }}"><input type="hidden" name="action" value="confirm"><input type="hidden" name="preview_token" value="{{ token }}">
{% for group in review %}<details class="border rounded mb-2" open><summary class="receipt-summary"><strong>{{ group.uploaded[0].serial_number }}</strong> · {{ 'New receipt' if group.kind == 'new' else 'Review existing serial / multiple events' }}</summary><div class="p-3"><label class="me-3"><input type="radio" name="choice_{{ loop.index0 }}" value="skip" {{ 'checked' if group.kind != 'new' else '' }}> Skip</label><label><input type="radio" name="choice_{{ loop.index0 }}" value="add" {{ 'checked' if group.kind == 'new' else '' }}> Add receipt event(s), keep history</label><div class="table-responsive"><table class="table table-sm"><thead><tr><th>Version</th>{% for field in fields %}<th>{{ field.replace('_',' ') }}</th>{% endfor %}</tr></thead><tbody>{% for row in group.existing %}<tr><th>Existing</th>{% for field in fields %}<td>{{ canonical_item(row[field]) if field == 'part_number' else (row[field] if row[field] is not none else '') }}</td>{% endfor %}</tr>{% endfor %}{% for row in group.uploaded %}<tr class="table-success"><th>To add</th>{% for field in fields %}<td>{{ row[field] }}</td>{% endfor %}</tr>{% endfor %}</tbody></table></div></div></details>{% endfor %}
<p class="small text-muted">Confirm within 30 minutes. POD lines and affected receipt history are checked again under database locks. Changes invalidate the preview; exact retries cannot duplicate receipts.</p><button class="btn btn-success">Confirm receiving</button></form></div></section>{% endif %}
<section class="card-lite bg-white receipt-card"><div class="card-header"><strong>Receiving history · latest 500 receipts</strong></div><div class="receipt-history"><table class="table table-sm"><thead><tr><th>Date</th><th>Invoice #</th><th>Box #</th><th>Part</th><th>Serial</th><th>Qty</th></tr></thead><tbody>{% for row in receipts %}<tr><td>{{ row.entry_date }}</td><td>{{ row.invoice_number or '' }}</td><td>{{ row.box_number or '' }}</td><td>{{ canonical_item(row.part_number) }}</td><td>{{ row.serial_number }}</td><td>{{ display_value(row.quantity) }}</td></tr>{% else %}<tr><td colspan="6" class="text-muted">No receiving records for this POD.</td></tr>{% endfor %}</tbody></table></div></section>
</main><template id="receipt-line-template"><tr><td><select class="form-select" name="part_number" required aria-label="Part">{% for part in parts %}<option value="{{ part }}">{{ part }}</option>{% endfor %}</select></td><td><textarea class="form-control" name="serial_number" rows="2" required aria-label="Serial numbers" placeholder="One serial per line, or NA"></textarea></td><td><input class="form-control" name="quantity" type="number" value="1" min="0.000001" step="any" required aria-label="QTY"></td><td><button type="button" class="btn btn-outline-danger" data-remove-receipt aria-label="Remove receipt row">Remove</button></td></tr></template>
<noscript><p class="alert alert-info">Enable JavaScript for manual receiving rows, or use the workbook upload above.</p></noscript>
<script type="application/json" id="receiving-initial">{{ manual|tojson }}</script>
<script src="/static/purchase-receiving.js?v=2" defer></script></body></html>'''
