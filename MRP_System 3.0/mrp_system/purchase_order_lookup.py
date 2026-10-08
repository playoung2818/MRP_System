"""Read-only, paginated purchase-order search."""

from datetime import date, datetime
from io import BytesIO
from urllib.parse import urlencode

from flask import Blueprint, current_app, render_template_string, request, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from mrp_system.runtime.db_config import get_engine

purchase_orders = Blueprint("purchase_orders", __name__)
PAGE_SIZE = 100
COLUMNS = [
    "QB Num", "Source Name", "Item", "Deliv Date", "Ship Date", "Order Date",
    "Qty", "Rcv'd", "Qty(+)", "Amount", "Inventory Site", "Name", "POD#", "Memo",
    "POD Ship Date Raw", "Shipping Schedule Ship Date", "Shipping Schedule Qty(+)",
    "Shipping Schedule Order Qty", "Shipping Schedule Item", "Shipping Schedule Reference",
]


def build_query(args):
    filters = []
    params = {}
    for field, column in [("qb_num", "QB Num"), ("source_name", "Source Name")]:
        value = args.get(field, "").strip()
        if value:
            # Literal substring matching: %, _ and backslash are not wildcards.
            params[field] = "%" + value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
            filters.append(f'"{column}" ILIKE :{field}')
    for field, column in [("deliv_date", "Deliv Date"), ("ship_date", "Ship Date"), ("source_names", "Source Name")]:
        values = args.get(field, [])
        if isinstance(values, str):
            values = [values] if values else []
        parts = []
        for index, value in enumerate(dict.fromkeys(values)):
            if value == "__blank__":
                parts.append('coalesce(btrim("Source Name"), \'\') = \'\'' if field == "source_names" else f'"{column}" IS NULL')
            elif value == "__none__":
                parts.append("FALSE")
            elif field == "source_names":
                params[f"{field}_{index}"] = value
                parts.append(f'btrim("{column}") = :{field}_{index}')
            else:
                try:
                    params[f"{field}_{index}"] = date.fromisoformat(value)
                except ValueError:
                    raise ValueError(f"Enter a valid {column} in YYYY-MM-DD format.") from None
                parts.append(f'"{column}"::date = :{field}_{index}')
        if parts:
            filters.append("(" + " OR ".join(parts) + ")")
    sort = args.get("sort", "deliv_date")
    direction = args.get("direction", "asc")
    if sort not in {"deliv_date", "ship_date", "source_names"} or direction not in {"asc", "desc"}:
        raise ValueError("Choose Delivery Date, Ship Date, or Source Name and a valid sort direction.")
    try:
        page = max(1, int(args.get("page", "1")))
    except ValueError:
        raise ValueError("Page must be a number.") from None
    where = " WHERE " + " AND ".join(filters) if filters else ""
    sort_col = {"deliv_date": "Deliv Date", "ship_date": "Ship Date", "source_names": "Source Name"}[sort]
    return where, params, sort_col, direction, page


def display_value(value):
    if value is None:
        return "—"
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value)


def request_filters():
    """Read the same filters for the paginated view and full Excel export."""
    args = {key: request.args.get(key, default).strip() for key, default in [
        ("qb_num", ""), ("source_name", ""), ("deliv_date", ""), ("ship_date", ""),
        ("sort", "deliv_date"), ("direction", "asc"), ("page", "1"),
    ]}
    for field in ("deliv_date", "ship_date", "source_names"):
        args[field] = [value.strip() for value in request.args.getlist(field) if value.strip()]
    return args


def select_sql(where, sort_col, direction):
    """Shared columns and stable ordering; sort choices are validated upstream."""
    fields = ", ".join('"' + col + '"' for col in COLUMNS)
    return f'''SELECT {fields} FROM public."Open_Purchase_Orders"{where}
        ORDER BY "{sort_col}" {direction.upper()} NULLS LAST,
            "QB Num" NULLS LAST, "Item" NULLS LAST,
            "Source Name" NULLS LAST, ctid'''


@purchase_orders.route("/purchase_orders")
def lookup():
    args = request_filters()
    filter_options = {"deliv_date": [], "ship_date": [], "source_names": []}
    rows, total, page, pages, error, status = [], 0, 1, 1, None, 200
    try:
        sort_choice = request.args.get("sort_choice")
        if sort_choice is not None:
            if sort_choice not in {f"{field}:{direction}" for field in ("deliv_date", "ship_date", "source_names") for direction in ("asc", "desc")}:
                raise ValueError("Choose a valid column sort.")
            args["sort"], args["direction"] = sort_choice.split(":")
            args["page"] = "1"
        where, params, sort_col, direction, page = build_query(args)
        engine = get_engine()
        try:
            with engine.connect() as conn:
                with conn.begin():
                    conn.execute(text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY"))
                    for field, column in [("deliv_date", "Deliv Date"), ("ship_date", "Ship Date"), ("source_names", "Source Name")]:
                        # Match the other filters, without restricting this column's choices.
                        option_where, option_params, *_ = build_query({**args, field: []})
                        expression = f'nullif(btrim("{column}"), \'\')' if field == "source_names" else f'"{column}"::date'
                        values = conn.execute(text(
                            f'SELECT DISTINCT {expression} AS value FROM public."Open_Purchase_Orders"'
                            + option_where + ' ORDER BY value NULLS LAST'
                        ), option_params).scalars().all()
                        filter_options[field] = [
                            {"value": (str(value) if field == "source_names" else value.isoformat()) if value is not None else "__blank__",
                             "label": (str(value) if field == "source_names" else value.isoformat()) if value is not None else "(Blanks)"}
                            for value in values
                        ]
                        existing = {option["value"] for option in filter_options[field]}
                        for selected in args[field]:
                            if selected not in existing and selected != "__none__":
                                filter_options[field].append({"value": selected, "label": "(Blanks)" if selected == "__blank__" else selected})
                    total = conn.execute(text('SELECT COUNT(*) FROM public."Open_Purchase_Orders"' + where), params).scalar_one()
                    pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
                    page = min(page, pages)
                    sql = select_sql(where, sort_col, direction) + " LIMIT :limit OFFSET :offset"
                    raw = conn.execute(text(sql), {**params, "limit": PAGE_SIZE, "offset": (page - 1) * PAGE_SIZE}).mappings()
                    rows = [{col: display_value(row[col]) for col in COLUMNS} for row in raw]
        finally:
            engine.dispose()
    except ValueError as exc:
        error, status = str(exc), 400
    except Exception:
        current_app.logger.exception("Purchase order lookup failed")
        error, status = "Purchase orders could not be loaded. Please try again.", 503

    def page_url(number):
        return "/purchase_orders?" + urlencode({**args, "page": number}, doseq=True)

    return render_template_string(
        TEMPLATE, filters=args, filter_options=filter_options, rows=rows, columns=COLUMNS, total=total, page=page, pages=pages,
        previous=page_url(page - 1), following=page_url(page + 1), error=error,
        first=(page - 1) * PAGE_SIZE + 1 if rows else 0,
        last=(page - 1) * PAGE_SIZE + len(rows) if rows else 0,
        loaded_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        numeric={"Qty", "Rcv'd", "Qty(+)", "Amount", "Shipping Schedule Qty(+)", "Shipping Schedule Order Qty"},
    ), status


def excel_cell(sheet, value):
    """Keep numbers/dates typed and prevent text from becoming Excel formulas."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        value = value.replace(tzinfo=None)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    elif isinstance(value, (date, datetime)):
        cell.number_format = "yyyy-mm-dd"
    return cell


@purchase_orders.route("/purchase_orders/export")
def export_excel():
    try:
        # Export all matches, regardless of the current results page.
        args = {**request_filters(), "page": "1"}
        where, params, sort_col, direction, _ = build_query(args)
    except ValueError as exc:
        return str(exc), 400

    workbook = None
    try:
        engine = get_engine()
        try:
            with engine.connect() as conn:
                with conn.begin():
                    conn.execute(text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY"))
                    total = conn.execute(text('SELECT COUNT(*) FROM public."Open_Purchase_Orders"' + where), params).scalar_one()
                    if total > 1_048_575:
                        return "Too many rows for one Excel worksheet. Narrow the purchase-order filters and try again.", 413
                    workbook = Workbook(write_only=True)
                    sheet = workbook.create_sheet("Purchase Orders")
                    sheet.freeze_panes = "A2"
                    sheet.append([excel_cell(sheet, column) for column in COLUMNS])
                    raw = conn.execution_options(yield_per=1000).execute(
                        text(select_sql(where, sort_col, direction)), params,
                    ).mappings()
                    for row in raw:
                        sheet.append([excel_cell(sheet, row[column]) for column in COLUMNS])
                    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{total + 1}"
        finally:
            engine.dispose()
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output, as_attachment=True,
            download_name=f"purchase_orders_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            max_age=0,
        )
    except Exception:
        current_app.logger.exception("Purchase order Excel export failed")
        return "Purchase orders could not be exported. Please try again.", 503
    finally:
        if workbook is not None:
            workbook.close()


TEMPLATE = """
<!doctype html><html lang="en"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Purchase Orders</title><link rel="icon" href="/static/robot.svg?v=1" type="image/svg+xml">
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<link href="{{ url_for('static', filename='workspace.css') }}?v=2" rel="stylesheet">
<style>
.po-filters{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}
 .po-heading{display:flex;align-items:center;justify-content:space-between;gap:10px}
.po-filter-toggle{border:1px solid #b6c4d6;border-radius:2px;background:#fff;color:#526479;padding:0 4px;line-height:20px;min-width:22px;font-size:12px}
.po-filter-toggle.is-filtered{background:#246bc3;color:#fff;border-color:#246bc3}
.po-filter-menu{position:fixed;z-index:2100;width:290px;max-width:calc(100vw - 20px);max-height:calc(100vh - 20px);overflow:auto;background:white;border:1px solid #bec9d6;border-radius:4px;box-shadow:0 8px 28px #24364c33;padding:6px 0;color:#253346;font-size:14px}
.po-filter-menu[hidden]{display:none}
.po-filter-menu [hidden]{display:none!important}
.po-menu-title{font-size:12px;font-weight:650;padding:6px 12px;color:#65768a;margin:0}
.po-menu-command{display:block;width:100%;text-align:left;border:0;background:white;padding:8px 12px;color:inherit;font-size:14px}
.po-menu-command:hover,.po-menu-command:focus-visible{background:#edf4fd}
.po-menu-command[aria-pressed="true"]{color:#246bc3;font-weight:650}
.po-menu-command:disabled{color:#99a3af;cursor:default}
.po-menu-section{border-top:1px solid #e0e5eb;margin-top:5px;padding:10px 12px}
.po-date-list{max-height:190px;overflow:auto;margin-top:8px;border:1px solid #e0e5eb;padding:5px}
.po-date-list label,.po-select-all{display:flex;align-items:center;gap:8px;padding:4px;font-size:13px;cursor:pointer}
.po-date-list input,.po-select-all input{accent-color:#246bc3;width:15px;height:15px}
.po-menu-actions{display:flex;justify-content:flex-end;gap:8px;padding:0 12px 6px}
.po-actions{grid-column:1/-1;display:flex;gap:10px;align-items:end;flex-wrap:wrap}
.po-actions label{display:block}.po-actions select{min-height:38px;font-size:14px}
.po-table{white-space:nowrap}.po-table .po-long{white-space:normal;min-width:240px;max-width:400px}
.po-pages{display:flex;align-items:center;justify-content:space-between;padding:12px 14px;gap:12px;flex-wrap:wrap}
.po-pages a{font-size:14px}.po-note{font-size:12px;color:var(--muted);margin:8px 0 16px}
@media(max-width:900px){.po-filters{grid-template-columns:repeat(2,minmax(0,1fr))}}
@media(max-width:540px){.po-filters{grid-template-columns:1fr}}
</style></head><body class="workspace-page">
<a class="workspace-skip" href="#workspace-content">Skip to content</a>
<header class="workspace-nav"><div class="workspace-nav-inner">
<a class="workspace-brand" href="/"><span class="workspace-brand-mark"><svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="icon icon-tabler icons-tabler-outline icon-tabler-robot" aria-hidden="true" focusable="false"><path stroke="none" d="M0 0h24v24H0z" fill="none" /><path d="M6 6a2 2 0 0 1 2 -2h8a2 2 0 0 1 2 2v4a2 2 0 0 1 -2 2h-8a2 2 0 0 1 -2 -2l0 -4" /><path d="M12 2v2" /><path d="M9 12v9" /><path d="M15 12v9" /><path d="M5 16l4 -2" /><path d="M15 14l4 2" /><path d="M9 18h6" /><path d="M10 8v.01" /><path d="M14 8v.01" /></svg></span><span><strong>NEOUSYS</strong><small>MANUFACTURING WORKSPACE</small></span></a>
<a class="workspace-home" href="/">Home</a></div></header>
<main class="workspace-main" id="workspace-content">
<div class="workspace-breadcrumb"><a href="/">Workspace</a><span>/</span><span>Purchase Orders</span></div>
<header class="workspace-header"><div><h1 class="workspace-title">Purchase orders</h1>
<p class="workspace-description">Find open purchase-order details and review delivery and shipping dates.</p></div>
<span class="workspace-loaded">Loaded {{ loaded_at }}</span></header>
<form id="po-search" class="workspace-search po-filters" method="get">
<input type="hidden" name="sort" value="{{ filters.sort }}">
<input type="hidden" name="direction" value="{{ filters.direction }}">
{% for field in ['deliv_date','ship_date','source_names'] %}{% for value in filters[field] %}<input type="hidden" name="{{ field }}" value="{{ value }}">{% endfor %}{% endfor %}
{% for field, label, kind in [('qb_num','QB Num','text'),('source_name','Source Name','text')] %}
<div><label class="form-label" for="{{ field }}">{{ label }}</label><input class="form-control" id="{{ field }}" name="{{ field }}" type="{{ kind }}" value="{{ filters[field] }}" {% if kind == 'text' %}placeholder="Contains…"{% endif %}></div>
{% endfor %}
<div class="po-actions">
<button class="btn btn-primary" type="submit">Search</button><a class="btn btn-outline-secondary" href="/purchase_orders">Clear filters</a>
<button class="btn btn-outline-success" type="submit" formaction="{{ url_for('purchase_orders.export_excel') }}" title="Download all rows matching these filters">Export Excel</button></div>
</form>
<p class="po-note">Open the Source Name or date column dropdown to sort or select values. All filters can apply together. Sorting uses one column at a time. The date 2099-12-31 represents a pending date.</p>
{% if error %}<div class="alert alert-warning" role="alert">{{ error }}</div>{% endif %}
<section class="card-lite bg-white"><div class="card-header"><strong>Open purchase orders</strong><span class="workspace-table-note">{{ total }} matching rows</span></div>
<div class="card-body"><div class="table-responsive"><table class="table table-sm table-bordered table-hover align-middle po-table">
<thead><tr>{% for col in columns %}
{% set date_field = 'deliv_date' if col == 'Deliv Date' else 'ship_date' if col == 'Ship Date' else 'source_names' if col == 'Source Name' else '' %}
<th scope="col" class="{{ 'text-end' if col in numeric else '' }}" {% if date_field %}aria-sort="{{ ('ascending' if filters.direction == 'asc' else 'descending') if filters.sort == date_field else 'none' }}"{% endif %}>
<div class="po-heading"><span>{{ 'Delivery Date' if col == 'Deliv Date' else col }}{% if date_field and filters.sort == date_field %} {{ '?' if filters.direction == 'asc' else '?' }}{% endif %}</span>
{% if date_field %}<button type="button" class="po-filter-toggle {{ 'is-filtered' if filters[date_field] else '' }}" data-open-filter="{{ date_field }}" aria-haspopup="dialog" aria-expanded="false" aria-controls="menu-{{ date_field }}" aria-label="Filter and sort {{ 'Delivery Date' if col == 'Deliv Date' else col }}{{ ', filter active' if filters[date_field] else '' }}">?</button>{% endif %}</div></th>{% endfor %}</tr></thead>
<tbody>{% for row in rows %}<tr>{% for col in columns %}<td class="{{ 'text-end' if col in numeric else '' }} {{ 'po-long' if col in ['Memo','Shipping Schedule Item','Shipping Schedule Reference'] else '' }}">{{ row[col] }}</td>{% endfor %}</tr>
{% else %}<tr><td colspan="{{ columns|length }}" class="text-center text-muted py-4">{{ 'Results unavailable.' if error else 'No purchase orders match these filters.' }}</td></tr>{% endfor %}</tbody>
</table></div></div><div class="po-pages"><span class="text-muted">Showing {{ first }}–{{ last }} of {{ total }}</span>
<div class="d-flex align-items-center gap-3">{% if page > 1 %}<a href="{{ previous }}">Previous</a>{% endif %}<span>Page {{ page }} of {{ pages }}</span>{% if page < pages %}<a href="{{ following }}">Next</a>{% endif %}</div></div></section>
</main>
{% for field, label in [('deliv_date','Delivery Date'),('ship_date','Ship Date'),('source_names','Source Name')] %}
<div id="menu-{{ field }}" class="po-filter-menu" role="dialog" aria-labelledby="title-{{ field }}" hidden data-field="{{ field }}">
<h2 class="po-menu-title" id="title-{{ field }}">{{ label }}</h2>
<button type="button" class="po-menu-command" data-sort="asc" aria-pressed="{{ 'true' if filters.sort == field and filters.direction == 'asc' else 'false' }}">&#9650; &nbsp; {{ 'Sort A to Z' if field == 'source_names' else 'Sort oldest to newest' }}</button>
<button type="button" class="po-menu-command" data-sort="desc" aria-pressed="{{ 'true' if filters.sort == field and filters.direction == 'desc' else 'false' }}">&#9660; &nbsp; {{ 'Sort Z to A' if field == 'source_names' else 'Sort newest to oldest' }}</button>
<button type="button" class="po-menu-command" data-clear {% if not filters[field] %}disabled{% endif %}>Clear filter from ?{{ label }}?</button>
<div class="po-menu-section">
<label class="visually-hidden" for="search-{{ field }}">Search {{ label }} values</label><input type="search" class="form-control po-date-search" id="search-{{ field }}" placeholder="{{ 'Search source names…' if field == 'source_names' else 'Search dates…' }}" autocomplete="off">
<label class="po-select-all"><input type="checkbox" data-select-all> <span>Select all</span></label>
<div class="po-date-list">
{% for option in filter_options[field] %}<label data-date-option><input type="checkbox" value="{{ option.value }}" {% if not filters[field] or option.value in filters[field] %}checked{% endif %}> <span>{{ option.label }}</span></label>{% endfor %}
<p class="text-muted small p-2 po-no-dates" hidden>No matching values.</p>
</div></div>
<div class="po-menu-actions"><button type="button" class="btn btn-primary" data-apply>OK</button><button type="button" class="btn btn-outline-secondary" data-cancel>Cancel</button></div></div>
{% endfor %}
<noscript><p class="alert alert-info">Enable JavaScript to use the column filter dropdowns.</p></noscript>
<script>
(() => {
  const form = document.getElementById('po-search');
  let menu = null, opener = null;
  function close(restore = true) {
    if (!menu) return;
    menu.hidden = true; opener.setAttribute('aria-expanded', 'false');
    if (restore) opener.focus();
    menu = null;
  }
  function choices(panel) { return [...panel.querySelectorAll('[data-date-option]')]; }
  function sync(panel) {
    const visible = choices(panel).filter(row => !row.hidden);
    const count = visible.filter(row => row.querySelector('input').checked).length;
    const all = panel.querySelector('[data-select-all]');
    all.checked = visible.length > 0 && count === visible.length;
    all.indeterminate = count > 0 && count < visible.length;
    all.disabled = !visible.length;
    panel.querySelector('.po-no-dates').hidden = visible.length > 0;
  }
  function clearField(field) { [...form.elements].filter(input => input.name === field).forEach(input => input.remove()); }
  function submit() { form.requestSubmit(); }
  document.querySelectorAll('[data-open-filter]').forEach(button => button.addEventListener('click', () => {
    const next = document.getElementById(button.getAttribute('aria-controls'));
    if (menu === next) { close(); return; }
    close(false); menu = next; opener = button;
    menu.querySelector('.po-date-search').value = '';
    choices(menu).forEach(row => { row.hidden = false; const input = row.querySelector('input'); input.checked = input.defaultChecked; });
    sync(menu); menu.hidden = false; button.setAttribute('aria-expanded', 'true');
    const rect = button.getBoundingClientRect();
    menu.style.left = Math.max(10, Math.min(rect.right - menu.offsetWidth, window.innerWidth - menu.offsetWidth - 10)) + 'px';
    menu.style.top = Math.max(10, Math.min(rect.bottom + 4, window.innerHeight - menu.offsetHeight - 10)) + 'px';
    menu.querySelector('[data-sort]').focus();
  }));
  document.querySelectorAll('.po-filter-menu').forEach(panel => {
    const field = panel.dataset.field;
    panel.querySelector('.po-date-search').addEventListener('input', event => {
      const query = event.target.value.trim().toLowerCase();
      choices(panel).forEach(row => { row.hidden = !row.textContent.toLowerCase().includes(query); });
      sync(panel);
    });
    panel.querySelector('[data-select-all]').addEventListener('change', event => {
      choices(panel).filter(row => !row.hidden).forEach(row => { row.querySelector('input').checked = event.target.checked; });
      sync(panel);
    });
    panel.querySelector('.po-date-list').addEventListener('change', () => sync(panel));
    panel.querySelectorAll('[data-sort]').forEach(button => button.addEventListener('click', () => {
      form.elements.namedItem('sort').value = field;
      form.elements.namedItem('direction').value = button.dataset.sort;
      submit();
    }));
    panel.querySelector('[data-clear]').addEventListener('click', () => { clearField(field); submit(); });
    panel.querySelector('[data-apply]').addEventListener('click', () => {
      const all = choices(panel).map(row => row.querySelector('input'));
      const selected = all.filter(input => input.checked).map(input => input.value);
      clearField(field);
      const values = selected.length === all.length && all.length ? [] : selected.length ? selected : ['__none__'];
      values.forEach(value => { const input = document.createElement('input'); input.type = 'hidden'; input.name = field; input.value = value; form.append(input); });
      submit();
    });
    panel.querySelector('[data-cancel]').addEventListener('click', () => close());
  });
  document.addEventListener('click', event => { if (menu && !menu.contains(event.target) && !opener.contains(event.target)) close(false); });
  document.addEventListener('keydown', event => { if (menu && event.key === 'Escape') { event.preventDefault(); close(); } });
  document.addEventListener('scroll', event => { if (menu && !menu.contains(event.target)) close(false); }, true);
  window.addEventListener('resize', () => close(false));
})();
</script>
</body></html>
"""
