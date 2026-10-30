"""Receiving tests use mocks only: no database/table creation or live receipts."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import re
from unittest.mock import MagicMock, patch

from flask import Flask
from openpyxl import Workbook
import pytest
from sqlalchemy import Table, MetaData, Column, Integer, String, Date, Float
from werkzeug.datastructures import MultiDict

from mrp_system import purchase_receiving as receiving
from mrp_system import purchase_order_lookup as po

POD='POD-261003'
PARTS={'abc':'Abc'}


@pytest.fixture(autouse=True)
def pure_item_normalization():
    with patch.object(receiving,'normalize_item',side_effect=lambda value: receiving.clean(value).upper()):
        yield


def receipt(**extra):
    row=dict(entry_date='2026-10-29',invoice_number='INV',box_number='1',
             pod_number=POD,part_number='Abc',serial_number='SN-1',quantity=1.0)
    row.update(extra)
    return row


def table():
    return Table('receiving_log',MetaData(),Column('id',Integer),
                 *(Column(key,Date if key=='entry_date' else Float if key=='quantity' else String)
                   for key in receiving.KEYS),Column('Reference',String))


def engine():
    result=MagicMock()
    result.connect.return_value.__enter__.return_value.dialect.name='postgresql'
    return result


def test_pod_search_is_literal_and_combines_with_other_filters():
    where,params,*_=po.build_query({'pod_num':"2610%_\\'",'source_name':'Vendor'})
    assert '"POD#" ILIKE :pod_num OR "QB Num" ILIKE :pod_num' in where
    assert params['pod_num']=="%2610\\%\\_\\\\'%"
    assert "2610" not in where
    assert '"Source Name" ILIKE :source_name' in where
    assert receiving.pod_id('261003')==POD
    assert receiving.pod_id('pod # 261003')==POD
    with pytest.raises(ValueError):receiving.pod_id("261003'; DROP TABLE receiving_log")


def test_legacy_qb_filter_uses_the_single_pod_search():
    app=Flask(__name__)
    with app.test_request_context('/purchase_orders?qb_num=261003'):
        filters=po.request_filters()
    assert filters['pod_num']=='261003'
    assert 'qb_num' not in filters
    with app.test_request_context('/purchase_orders?pod_num=261004&qb_num=261003'):
        assert po.request_filters()['pod_num']=='261004'


def test_source_context_keeps_source_fingerprint_raw_and_canonicalizes_display():
    conn=MagicMock()
    raw=dict(Item='ABC',**{'POD#':POD,'Qty':5})
    source_result,master_result=MagicMock(),MagicMock()
    source_result.mappings.return_value=[raw]
    master_result.scalars.return_value=['Abc']
    conn.execute.side_effect=[source_result,master_result]
    rows,parts,version=receiving.source_context(conn,POD)
    assert rows[0]['display_item']=='Abc'
    assert parts['abc']=='Abc'
    assert 'display_item' not in raw
    assert version==receiving.fingerprint([raw])
    sql,params=conn.execute.call_args_list[0].args
    assert POD not in str(sql)
    assert params['pod']==POD and params['digits']=='261003'


def test_validation_canonicalizes_only_new_receipts_and_rejects_cross_pod():
    rows,skipped=receiving.validate_rows([receipt(part_number='ABC'),receipt(serial_number='nan')],POD,PARTS)
    assert rows[0]['part_number']=='Abc'
    assert skipped==1
    with pytest.raises(ValueError,match='another POD'):
        receiving.validate_rows([receipt(pod_number='POD-261004')],POD,PARTS)
    with pytest.raises(ValueError,match='not on this POD'):
        receiving.validate_rows([receipt(part_number='Other')],POD,PARTS)


@pytest.mark.parametrize('quantity',[0,-1,'nan','inf','=1+1'])
def test_bad_quantities_block_entire_batch(quantity):
    with pytest.raises(ValueError,match='Nothing saved'):
        receiving.validate_rows([receipt(),receipt(quantity=quantity)],POD,PARTS)


def test_manual_serial_lines_and_na_bulk_quantity():
    form=MultiDict([('entry_date','2026-10-29'),('invoice_number','INV'),('box_number','BOX'),
                    ('part_number','Abc'),('serial_number','SN1\nSN2'),('quantity','2'),
                    ('part_number','Abc'),('serial_number','NA'),('quantity','12')])
    rows,_=receiving.validate_rows(receiving.manual_rows(form,POD),POD,PARTS)
    assert [r['serial_number'] for r in rows]==['SN1','SN2','NA']
    assert rows[-1]['quantity']==12


@pytest.mark.parametrize('serials,qty,expected_rows,expected_units,warning',[
    ('SN1\nSN2',2,2,2,False),
    ('SN1\n\nSN2\n',3,2,2,True),
    ('NA',12,1,12,False),
    ('NA\nna',12,1,12,False),
])
def test_manual_total_qty_serial_count_warning_and_na_exception(serials,qty,expected_rows,expected_units,warning):
    warnings=[]
    form=MultiDict([('entry_date','2026-10-29'),('part_number','Abc'),
                    ('serial_number',serials),('quantity',str(qty))])
    rows=receiving.manual_rows(form,POD,warnings)
    assert len(rows)==expected_rows
    assert sum(row['quantity'] for row in rows)==expected_units
    assert bool(warnings)==warning


@pytest.mark.parametrize('part,serial,valid',[
    ('i7-14700','SERIAL_BATCH',True),('Core-i7','SERIAL_BATCH',True),
    ('I9-14900','SERIAL',False),('Core Ultra','SERIAL_',False),
    ('i7-14700','_BATCH',False),('i7-14700','NA',True),
    ('Dust-i7','SERIAL',True),('RAM','SERIAL',True),
])
def test_cpu_serial_and_batch_validation(part,serial,valid):
    if valid:receiving.validate_cpu_batches([receipt(part_number=part,serial_number=serial)])
    else:
        with pytest.raises(ValueError,match='both serial number and batch'):
            receiving.validate_cpu_batches([receipt(part_number=part,serial_number=serial)])


def test_workbook_uses_first_sheet_and_requires_full_headers():
    workbook=Workbook();sheet=workbook.active;sheet.title='Anything'
    sheet.append(list(receiving.HEADERS))
    sheet.append([date(2026,10,29),'INV','1',POD,'ABC','SN-1',None])
    workbook.create_sheet('Receiving');workbook.active=1
    content=BytesIO();workbook.save(content);workbook.close()
    rows,_=receiving.validate_rows(receiving.upload_rows(content.getvalue()),POD,PARTS)
    assert rows[0]==receipt()
    with pytest.raises(ValueError):receiving.upload_rows(b'not an Excel file')


def test_preview_skips_exact_repeats_and_preserves_distinct_na_events():
    conn=MagicMock()
    existing=dict(receipt(),id=1,Reference='Preserve')
    rows=[receipt(),receipt(serial_number='NA'),receipt(serial_number='NA'),
          receipt(serial_number='NA',invoice_number='OTHER')]
    with patch.object(receiving,'receiving_table',return_value=table()),patch.object(
            receiving,'database_groups',return_value={'SN-1':[existing]}):
        review,skipped=receiving.preview(conn,rows)
    assert len(review)==2
    assert skipped==2
    assert all(g['kind']=='new' for g in review)
    assert existing['Reference']=='Preserve'
    conn.commit.assert_not_called()


def apply(review,current=None,choices=None,source='source'):
    db=engine();conn=db.connect.return_value.__enter__.return_value
    with patch.object(receiving,'source_context',return_value=([],PARTS,source)),patch.object(
            receiving,'receiving_table',return_value=table()),patch.object(
            receiving,'database_groups',return_value=current or {}):
        result=receiving.apply_receipts(db,POD,'source',review,choices or {})
    return result,conn


def group(row=None,matches=None):
    row=row or receipt();matches=matches or []
    return dict(key=receiving.receipt_key(row),uploaded=[row],snapshot=receiving.fingerprint(matches))


def test_confirm_append_only_no_history_update_or_delete():
    original=dict(receipt(entry_date='2026-09-01'),id=1,Reference='Preserve reference')
    before=deepcopy(original)
    result,conn=apply([group(matches=[original])],{'SN-1':[original]},{'0':'add'})
    assert result==(1,0)
    sql=[str(call.args[0]) for call in conn.execute.call_args_list]
    assert any('LOCK TABLE public.receiving_log IN SHARE ROW EXCLUSIVE MODE' in s for s in sql)
    assert sum(s.startswith('INSERT INTO') for s in sql)==1
    assert not any(s.startswith(('UPDATE ','DELETE ','CREATE ','ALTER ')) for s in sql)
    assert original==before
    conn.commit.assert_called_once()
    conn.rollback.assert_not_called()


def test_exact_retry_is_idempotent_even_when_snapshot_changed():
    result,conn=apply([group()],{'SN-1':[dict(receipt(),id=1,Reference='Keep')]},{'0':'add'})
    assert result==(0,1)
    assert not any(str(c.args[0]).startswith('INSERT') for c in conn.execute.call_args_list)


def test_conflicts_default_to_skip():
    result,conn=apply([group()])
    assert result==(0,1)
    assert not any(str(c.args[0]).startswith('INSERT') for c in conn.execute.call_args_list)


@pytest.mark.parametrize('changed_source',[False,True])
def test_stale_receipt_or_source_rolls_back_before_insert(changed_source):
    db=engine();conn=db.connect.return_value.__enter__.return_value
    with (patch.object(receiving,'source_context',return_value=([],PARTS,'changed' if changed_source else 'source')),
          patch.object(receiving,'receiving_table',return_value=table()),patch.object(
          receiving,'database_groups',return_value={'SN-1':[dict(receipt(quantity=3),id=2)]})):
        with pytest.raises(ValueError,match='changed since preview'):
            receiving.apply_receipts(db,POD,'source',[group()],{'0':'add'})
    conn.rollback.assert_called_once()
    conn.commit.assert_not_called()
    assert not any(str(c.args[0]).startswith('INSERT') for c in conn.execute.call_args_list)


def test_insert_failure_rolls_back_the_entire_batch():
    db=engine();conn=db.connect.return_value.__enter__.return_value
    def execute(statement,*args):
        if str(statement).startswith('INSERT'):
            raise RuntimeError('simulated insert failure')
        return MagicMock()
    conn.execute.side_effect=execute
    with (patch.object(receiving,'source_context',return_value=([],PARTS,'source')),
          patch.object(receiving,'receiving_table',return_value=table()),
          patch.object(receiving,'database_groups',return_value={})):
        with pytest.raises(RuntimeError,match='simulated insert failure'):
            receiving.apply_receipts(db,POD,'source',[group()],{'0':'add'})
    conn.rollback.assert_called_once()
    conn.commit.assert_not_called()


@pytest.fixture
def client():
    app=Flask(__name__);app.secret_key='unit-test-only'
    app.register_blueprint(po.purchase_orders)
    db=engine();conn=db.connect.return_value.__enter__.return_value
    conn.execute.return_value.mappings.return_value=[]
    source=[dict(Item='ABC',display_item='Abc',**{'Source Name':receiving.WORKBOOK_VENDOR,'Qty':10,"Rcv'd":0,
                  'Qty(+)':10,'Inventory Site':'WH01S-NTA','Deliv Date':date(2026,11,1)})]
    with patch.object(receiving,'get_engine',return_value=db),patch.object(
            receiving,'source_context',return_value=(source,PARTS,'source')),patch.object(
            receiving,'receiving_table',return_value=table()),patch.object(receiving,'database_groups',return_value={}):
        yield app.test_client(),db


def manual_data(token):
    return dict(csrf_token=token,action='manual',entry_date='2026-10-29',invoice_number='INV',box_number='1',
                part_number='Abc',serial_number='SN-1',quantity='1')


def test_detail_get_and_preview_do_not_write_and_preserve_manual_values(client):
    browser,db=client
    response=browser.get('/purchase_orders/pod/'+POD)
    assert response.status_code==200
    assert 'Receive this POD' in response.text and 'Preview workbook' in response.text
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    response=browser.post('/purchase_orders/pod/'+POD,data=manual_data(csrf))
    assert response.status_code==200
    assert 'nothing saved yet' in response.text
    assert 'name="preview_token"' in response.text
    assert '"serial_number": "SN-1"' in response.text
    db.connect.return_value.__enter__.return_value.commit.assert_not_called()


@pytest.mark.parametrize('workbook_pod,expected',[(POD,200),('POD-261004',400)])
def test_workbook_route_previews_only_this_pod_and_never_writes(client,workbook_pod,expected):
    browser,db=client
    browser.get('/purchase_orders/pod/'+POD)
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    book=Workbook();sheet=book.active
    sheet.append(list(receiving.HEADERS))
    sheet.append([date(2026,10,29),'INV','1',workbook_pod,'ABC','SN-1',1])
    content=BytesIO();book.save(content);book.close();content.seek(0)
    response=browser.post('/purchase_orders/pod/'+POD,data=dict(
        csrf_token=csrf,action='upload',file=(content,'receiving.xlsx')))
    assert response.status_code==expected
    if expected==200:assert 'nothing saved yet' in response.text
    else:assert 'another POD' in response.text
    db.connect.return_value.__enter__.return_value.commit.assert_not_called()


@pytest.mark.parametrize('vendors',[['Other vendor'],[''],[receiving.WORKBOOK_VENDOR,'Other vendor']])
def test_other_vendors_have_manual_receiving_without_workbook_module(client,vendors):
    browser,_=client
    source=[dict(Item='ABC',display_item='Abc',**{'Source Name':vendor,'Qty':10,"Rcv'd":0,
                  'Qty(+)':10,'Inventory Site':'WH01S-NTA','Deliv Date':date(2026,11,1)}) for vendor in vendors]
    with patch.object(receiving,'source_context',return_value=(source,PARTS,'source')):
        response=browser.get('/purchase_orders/pod/'+POD)
        assert response.status_code==200
        assert 'Receive this POD' in response.text
        assert 'Or upload receiving workbook' not in response.text
        assert 'type="file"' not in response.text
        assert re.search(r'id="receipt-box"[^>]*value="1--1"[^>]*readonly',response.text)
        with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
        response=browser.post('/purchase_orders/pod/'+POD,data=dict(csrf_token=csrf,action='upload'))
        assert response.status_code==400
        assert 'Use manual receiving' in response.text
        response=browser.post('/purchase_orders/pod/'+POD,data=manual_data(csrf))
        assert response.status_code==200
        assert 'nothing saved yet' in response.text
        from itsdangerous import URLSafeTimedSerializer
        token=re.search(r'name="preview_token" value="([^"]+)"',response.text).group(1)
        payload=URLSafeTimedSerializer('unit-test-only',salt='pod-receiving-v1').loads(token)
        assert payload['review'][0]['uploaded'][0]['box_number']=='1--1'
        assert payload['review'][0]['uploaded'][0]['invoice_number']==receiving.vendor_label(source)
        assert re.search(r'id="receipt-invoice"[^>]*readonly',response.text)


@pytest.mark.parametrize('box',['1--1','tampered'])
def test_other_vendor_box_is_enforced_again_at_confirm(box):
    db=engine();conn=db.connect.return_value.__enter__.return_value
    with (patch.object(receiving,'source_context',return_value=([{'Source Name':'Other vendor'}],PARTS,'source')),
          patch.object(receiving,'receiving_table',return_value=table()),
          patch.object(receiving,'database_groups',return_value={})):
        if box=='1--1':
            assert receiving.apply_receipts(db,POD,'source',[group(receipt(box_number=box,invoice_number='Other vendor'))],{'0':'add'})==(1,0)
            conn.commit.assert_called_once()
        else:
            with pytest.raises(ValueError,match='Box # must be 1--1'):
                receiving.apply_receipts(db,POD,'source',[group(receipt(box_number=box))],{'0':'add'})
            conn.rollback.assert_called_once()
            conn.commit.assert_not_called()


@pytest.mark.parametrize('vendor,serial,status',[
    ('Other vendor','SERIAL',400),('Other vendor','SERIAL_BATCH',200),
    (receiving.WORKBOOK_VENDOR,'SERIAL',200),
])
def test_cpu_preview_requires_combined_batch_only_for_other_vendors(client,vendor,serial,status):
    browser,_=client
    source=[dict(Item='i7-14700',display_item='i7-14700',**{'Source Name':vendor,'Qty':10,"Rcv'd":0,
                'Qty(+)':10,'Inventory Site':'WH01S-NTA','Deliv Date':date(2026,11,1)})]
    with patch.object(receiving,'source_context',return_value=(source,{'i7-14700':'i7-14700'},'source')):
        browser.get('/purchase_orders/pod/'+POD)
        with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
        data=manual_data(csrf);data.update(part_number='i7-14700',serial_number=serial)
        response=browser.post('/purchase_orders/pod/'+POD,data=data)
        assert response.status_code==status
        if status==400:assert 'both serial number and batch number' in response.text
        else:assert 'nothing saved yet' in response.text


def test_preview_warns_on_manual_qty_mismatch_without_multiplying_receipts(client):
    from itsdangerous import URLSafeTimedSerializer
    browser,_=client
    browser.get('/purchase_orders/pod/'+POD)
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    data=manual_data(csrf);data.update(serial_number='SN1\nSN2',quantity='3')
    response=browser.post('/purchase_orders/pod/'+POD,data=data)
    assert response.status_code==200
    assert '<th>QTY</th>' in response.text
    assert '2 serial-number lines do not match QTY 3' in response.text
    token=re.search(r'name="preview_token" value="([^"]+)"',response.text).group(1)
    payload=URLSafeTimedSerializer('unit-test-only',salt='pod-receiving-v1').loads(token)
    assert sum(row['quantity'] for group in payload['review'] for row in group['uploaded'])==2
    data.update(serial_number='NA\nNA',quantity='12')
    response=browser.post('/purchase_orders/pod/'+POD,data=data)
    assert response.status_code==200
    assert 'class="alert alert-warning"' not in response.text


def test_csrf_and_tampered_token_do_not_call_writer(client):
    browser,_=client
    browser.get('/purchase_orders/pod/'+POD)
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    with patch.object(receiving,'apply_receipts') as writer:
        assert browser.post('/purchase_orders/pod/'+POD,data=manual_data('wrong')).status_code==400
        response=browser.post('/purchase_orders/pod/'+POD,data=dict(csrf_token=csrf,action='confirm',preview_token='bad'))
        assert response.status_code==400
        writer.assert_not_called()


def test_expired_preview_does_not_call_writer(client):
    from itsdangerous import URLSafeTimedSerializer
    browser,_=client
    browser.get('/purchase_orders/pod/'+POD)
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    with patch('itsdangerous.timed.TimestampSigner.get_timestamp',return_value=1):
        token=URLSafeTimedSerializer('unit-test-only',salt='pod-receiving-v1').dumps(
            dict(version=1,csrf=csrf,pod=POD,source='source',review=[group()]))
    with patch.object(receiving,'apply_receipts') as writer:
        response=browser.post('/purchase_orders/pod/'+POD,data=dict(csrf_token=csrf,action='confirm',preview_token=token))
        assert response.status_code==400
        assert 'expired or invalid' in response.text
        writer.assert_not_called()


def test_signed_preview_confirm_redirect_and_cross_pod_rejected(client):
    browser,_=client
    browser.get('/purchase_orders/pod/'+POD)
    with browser.session_transaction() as s:csrf=s['pod_receiving_csrf']
    response=browser.post('/purchase_orders/pod/'+POD,data=manual_data(csrf))
    token=re.search(r'name="preview_token" value="([^"]+)"',response.text).group(1)
    form=dict(csrf_token=csrf,action='confirm',preview_token=token,choice_0='add')
    with patch.object(receiving,'apply_receipts',return_value=(1,0)) as writer:
        response=browser.post('/purchase_orders/pod/'+POD,data=form)
        assert response.status_code==303
        assert writer.call_args.args[1]==POD
        writer.reset_mock()
        response=browser.post('/purchase_orders/pod/POD-261004',data=form)
        assert response.status_code==400
        writer.assert_not_called()
