from datetime import date
from io import BytesIO
from unittest.mock import MagicMock, patch

from flask import Flask
from openpyxl import load_workbook

from mrp_system import purchase_order_lookup as po


def export_request(rows, query="", total=None):
    app = Flask(__name__)
    app.register_blueprint(po.purchase_orders)
    engine = MagicMock()
    conn = engine.connect.return_value.__enter__.return_value
    conn.execution_options.return_value = conn
    queries = []

    def execute(statement, params=None):
        sql = str(statement)
        queries.append((sql, params or {}))
        result = MagicMock()
        if "COUNT(*)" in sql:
            result.scalar_one.return_value = len(rows) if total is None else total
        else:
            result.mappings.return_value = rows
        return result

    conn.execute.side_effect = execute
    with patch.object(po, "get_engine", return_value=engine):
        response = app.test_client().get("/purchase_orders/export" + query)
    engine.dispose.assert_called_once()
    return response, queries


def test_export_all_filtered_rows_preserves_types_and_text():
    rows = []
    for index in range(105):
        row = dict.fromkeys(po.COLUMNS)
        row.update({"QB Num": f"POD-{index:03}", "Source Name": "=HYPERLINK(\"evil\")",
                    "Deliv Date": date(2026, 10, 12), "Qty": 50, "Amount": 123.45,
                    "Memo": "@text\x01 with control character"})
        rows.append(row)
    response, queries = export_request(rows, "?page=2&source_names=Vendor+A&source_names=Vendor+B&deliv_date=2026-10-12&sort=source_names&direction=desc")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["Content-Disposition"]
    assert ".xlsx" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["Purchase Orders"]
    assert sheet.max_row == 106  # Not limited to the 100-row web page.
    assert [cell.value for cell in sheet[1]] == po.COLUMNS
    assert sheet["B2"].value == '=HYPERLINK("evil")'
    assert sheet["B2"].data_type == "s"
    assert sheet["D2"].value.date() == date(2026, 10, 12)
    assert sheet["G2"].value == 50 and sheet["G2"].data_type == "n"
    assert sheet["J2"].value == 123.45
    assert sheet["N2"].value == "@text with control character"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:T106"
    sql, params = queries[-1]
    assert "LIMIT" not in sql and "OFFSET" not in sql
    assert 'ORDER BY "Source Name" DESC' in sql
    assert params["source_names_0"] == "Vendor A"
    assert params["source_names_1"] == "Vendor B"
    assert params["deliv_date_0"] == date(2026, 10, 12)
    workbook.close()


def test_empty_export_has_column_headers():
    response, _ = export_request([])
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert list(workbook.active.values) == [tuple(po.COLUMNS)]
    workbook.close()


def test_invalid_filters_rejected_before_database_access():
    app = Flask(__name__)
    app.register_blueprint(po.purchase_orders)
    with patch.object(po, "get_engine") as engine:
        response = app.test_client().get("/purchase_orders/export?deliv_date=bad-date")
    assert response.status_code == 400
    engine.assert_not_called()


def test_export_rejects_excel_row_limit():
    response, queries = export_request([], total=1_048_576)
    assert response.status_code == 413
    assert "Narrow" in response.get_data(as_text=True)
    assert "COUNT(*)" in queries[-1][0]


def test_export_database_error_returns_friendly_message():
    app = Flask(__name__)
    app.register_blueprint(po.purchase_orders)
    with patch.object(po, "get_engine", side_effect=RuntimeError("private DSN")):
        response = app.test_client().get("/purchase_orders/export")
    assert response.status_code == 503
    assert b"private DSN" not in response.data
